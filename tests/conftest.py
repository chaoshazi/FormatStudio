from __future__ import annotations

import os
import tempfile
from pathlib import Path

import pytest

TEST_DATA = Path(tempfile.mkdtemp(prefix="format-studio-tests-"))
os.environ["FORMAT_STUDIO_DATA_DIR"] = str(TEST_DATA)
os.environ["FORMAT_STUDIO_OUTPUT_DIR"] = str(TEST_DATA / "output")
os.environ["FORMAT_STUDIO_DISABLE_TOKEN"] = "1"

from fastapi.testclient import TestClient  # noqa: E402

from app import db  # noqa: E402
from app.main import app  # noqa: E402


@pytest.fixture(scope="session", autouse=True)
def initialized_db():
    db.init_db()


@pytest.fixture()
def client():
    with TestClient(app) as test_client:
        yield test_client


@pytest.fixture()
def work_dir(tmp_path: Path) -> Path:
    folder = tmp_path / "work"
    folder.mkdir()
    return folder


@pytest.fixture()
def txt_file(work_dir: Path) -> Path:
    path = work_dir / "示例.txt"
    path.write_text("第一段内容\n\n第二段内容\n", encoding="utf-8")
    return path


@pytest.fixture()
def md_file(work_dir: Path) -> Path:
    path = work_dir / "说明.md"
    path.write_text(
        "# 标题\n\n正文 **加粗** 与 *斜体*。\n\n- 项目一\n- 项目二\n\n| 名称 | 数量 |\n| --- | --- |\n| A | 2 |\n",
        encoding="utf-8",
    )
    return path


@pytest.fixture()
def csv_file(work_dir: Path) -> Path:
    path = work_dir / "数据.csv"
    path.write_text("姓名,部门,编号\n张三,研发,001\n李四,市场,002\n", encoding="utf-8-sig")
    return path


@pytest.fixture()
def docx_file(work_dir: Path) -> Path:
    from docx import Document

    document = Document()
    document.add_heading("示例文档", level=1)
    document.add_paragraph("这是正文段落，包含中文内容。")
    table = document.add_table(rows=3, cols=2)
    table.style = "Table Grid"
    for index, value in enumerate(["项目", "数量"]):
        table.cell(0, index).text = value
    table.cell(1, 0).text = "A"
    table.cell(1, 1).text = "2"
    table.cell(2, 0).text = "B"
    table.cell(2, 1).text = "3"
    path = work_dir / "示例.docx"
    document.save(str(path))
    return path


@pytest.fixture()
def xlsx_file(work_dir: Path) -> Path:
    import openpyxl

    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "数据"
    worksheet.append(["姓名", "部门", "编号"])
    worksheet.append(["张三", "研发", "001"])
    worksheet.append(["李四", "市场", "002"])
    path = work_dir / "数据.xlsx"
    workbook.save(str(path))
    return path


@pytest.fixture()
def pptx_file(work_dir: Path) -> Path:
    from pptx import Presentation

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    slide.shapes.title.text = "示例演示"
    slide.placeholders[1].text = "正文内容"
    path = work_dir / "示例.pptx"
    presentation.save(str(path))
    return path


@pytest.fixture()
def pdf_file(work_dir: Path) -> Path:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfgen import canvas

    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    path = work_dir / "示例.pdf"
    pdf = canvas.Canvas(str(path), pagesize=A4)
    pdf.setFont("STSong-Light", 18)
    pdf.drawString(72, 760, "示例 PDF 标题")
    pdf.setFont("STSong-Light", 11)
    pdf.drawString(72, 720, "这是第一段正文，用于测试文本提取。")
    pdf.drawString(72, 700, "第二行内容与第一段属于同一段落。")
    pdf.drawString(72, 650, "这是第二段正文。")
    x, y = 72, 560
    widths = [120, 80, 100]
    rows = [["姓名", "数量", "备注"], ["张三", "2", "测试"], ["李四", "3", "示例"]]
    for row_index, row in enumerate(rows):
        current_x = x
        for col_index, value in enumerate(row):
            pdf.rect(current_x, y - row_index * 24, widths[col_index], 24, stroke=1, fill=0)
            pdf.drawString(current_x + 5, y - row_index * 24 + 7, value)
            current_x += widths[col_index]
    pdf.showPage()
    pdf.setFont("STSong-Light", 12)
    pdf.drawString(72, 760, "第二页内容")
    pdf.save()
    return path