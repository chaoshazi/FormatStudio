import time
from pathlib import Path

from openpyxl import load_workbook


def test_health_and_capabilities(client):
    health = client.get("/api/v1/health")
    assert health.status_code == 200
    payload = health.json()
    assert payload["ok"] is True
    caps = client.get("/api/v1/capabilities").json()
    assert len(caps["formats"]) == 8
    assert len(caps["matrix"]) == 56


def test_txt_to_md_task(client, txt_file: Path):
    with txt_file.open("rb") as handle:
        response = client.post(
            "/api/v1/tasks",
            files=[("files", (txt_file.name, handle, "text/plain"))],
            data={"target_format": "md", "mode": "default"},
        )
    assert response.status_code == 202, response.text
    task_id = response.json()["id"]
    task = None
    for _ in range(60):
        task = client.get(f"/api/v1/tasks/{task_id}").json()
        if task["status"] in {"succeeded", "failed", "canceled", "partial", "interrupted"}:
            break
        time.sleep(0.2)
    assert task is not None
    assert task["status"] == "succeeded", task
    item = task["items"][0]
    assert item["download_url"]
    assert item["download_name"].endswith(".md")
    download = client.get(item["download_url"])
    assert download.status_code == 200
    assert "第一段内容".encode("utf-8") in download.content


def test_unsupported_pair(client, docx_file: Path):
    with docx_file.open("rb") as handle:
        response = client.post(
            "/api/v1/tasks",
            files=[("files", (docx_file.name, handle, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"))],
            data={"target_format": "pptx"},
        )
    assert response.status_code == 422
    assert response.json()["detail"]["code"] == "UNSUPPORTED_PAIR"


def test_settings_roundtrip(client, tmp_path: Path, work_dir: Path):
    output = work_dir / "custom-output"
    response = client.put(
        "/api/v1/settings",
        json={"output_dir": str(output), "retention_days": 30, "csv_all_text": True},
    )
    assert response.status_code == 200
    payload = response.json()
    assert payload["retention_days"] == 30
    assert payload["csv_all_text"] is True
    assert output.exists()


def test_template_upload_and_generate(client, work_dir: Path):
    from docx import Document

    document = Document()
    document.add_paragraph("姓名：{{name}}")
    document.add_paragraph("部门：{{dept}}")
    template_path = work_dir / "模板.docx"
    document.save(str(template_path))
    with template_path.open("rb") as handle:
        upload = client.post("/api/v1/templates", files=[("file", (template_path.name, handle, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"))], data={"name": "测试模板"})
    assert upload.status_code == 201, upload.text
    record = upload.json()
    assert {field["name"] for field in record["fields"]} == {"name", "dept"}

    data_path = work_dir / "data.csv"
    data_path.write_text("name,dept\n张三,研发\n李四,市场\n", encoding="utf-8-sig")
    with data_path.open("rb") as handle:
        generated = client.post(
            f"/api/v1/templates/{record['id']}/generate",
            files=[("data_file", (data_path.name, handle, "text/csv"))],
        )
    assert generated.status_code == 200, generated.text
    payload = generated.json()
    assert payload["generated"] == 2
    assert payload["archive_url"]