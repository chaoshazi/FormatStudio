from pathlib import Path

import pytest
from docx import Document
from pptx import Presentation

from app.services.pdf_conv import (
    PdfConversionError,
    pdf_to_docx_editable,
    pdf_to_docx_layout,
    pdf_to_md,
    pdf_to_pptx_editable,
    pdf_to_txt,
    pdf_to_xlsx,
)


def test_pdf_to_txt(pdf_file: Path, work_dir: Path):
    target = work_dir / "out.txt"
    pdf_to_txt(pdf_file, target)
    content = target.read_text(encoding="utf-8")
    assert "示例 PDF 标题" in content
    assert "第二页内容" in content
    assert "\f" in content


def test_pdf_to_md(pdf_file: Path, work_dir: Path):
    target = work_dir / "out.md"
    pdf_to_md(pdf_file, target)
    content = target.read_text(encoding="utf-8")
    assert "<!-- page: 1 -->" in content
    assert "<!-- page: 2 -->" in content


def test_pdf_to_docx_editable(pdf_file: Path, work_dir: Path):
    target = work_dir / "out.docx"
    warnings = pdf_to_docx_editable(pdf_file, target)
    assert target.exists()
    assert isinstance(warnings, list)
    document = Document(str(target))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "示例 PDF 标题" in text


def test_pdf_to_docx_layout(pdf_file: Path, work_dir: Path):
    target = work_dir / "layout.docx"
    pdf_to_docx_layout(pdf_file, target)
    document = Document(str(target))
    assert len(document.sections) == 2
    assert document.sections[0].page_width is not None
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "示例 PDF 标题" in text


def test_pdf_to_pptx(pdf_file: Path, work_dir: Path):
    target = work_dir / "out.pptx"
    pdf_to_pptx_editable(pdf_file, target)
    presentation = Presentation(str(target))
    assert len(presentation.slides) == 2
    texts = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                texts.append(shape.text_frame.text)
    assert any("示例 PDF 标题" in text for text in texts)


def test_pdf_table_to_xlsx(pdf_file: Path, work_dir: Path):
    target = work_dir / "tables.xlsx"
    count = pdf_to_xlsx(pdf_file, target)
    assert count >= 1
    import openpyxl

    workbook = openpyxl.load_workbook(target)
    values = [cell.value for row in workbook.worksheets[0].iter_rows() for cell in row]
    assert "姓名" in values
    workbook.close()


def test_scanned_pdf_needs_ocr(work_dir: Path):
    from PIL import Image, ImageDraw
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    image_path = work_dir / "scan.png"
    image = Image.new("RGB", (800, 1000), "white")
    draw = ImageDraw.Draw(image)
    draw.text((100, 100), "SCANNED PAGE", fill="black")
    image.save(image_path)
    pdf_path = work_dir / "scan.pdf"
    pdf = canvas.Canvas(str(pdf_path), pagesize=A4)
    pdf.drawImage(str(image_path), 40, 400, width=500, height=620)
    pdf.showPage()
    pdf.save()
    with pytest.raises(PdfConversionError) as exc:
        pdf_to_docx_editable(pdf_path, work_dir / "scan.docx")
    assert exc.value.code == "NEEDS_OCR"

def test_encrypted_pdf_password(work_dir: Path):
    from pypdf import PdfReader, PdfWriter
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    plain = work_dir / "plain.pdf"
    pdf = canvas.Canvas(str(plain), pagesize=A4)
    pdf.drawString(72, 700, "secret")
    pdf.showPage()
    pdf.save()
    encrypted = work_dir / "encrypted.pdf"
    reader = PdfReader(str(plain))
    writer = PdfWriter()
    for page in reader.pages:
        writer.add_page(page)
    writer.encrypt("pw123")
    with encrypted.open("wb") as handle:
        writer.write(handle)
    with pytest.raises(PdfConversionError) as missing:
        pdf_to_txt(encrypted, work_dir / "missing.txt")
    assert missing.value.code == "PDF_PASSWORD_REQUIRED"
    with pytest.raises(PdfConversionError) as wrong:
        pdf_to_txt(encrypted, work_dir / "wrong.txt", "bad")
    assert wrong.value.code == "PDF_PASSWORD_INVALID"
    pdf_to_txt(encrypted, work_dir / "ok.txt", "pw123")
    assert "secret" in (work_dir / "ok.txt").read_text(encoding="utf-8")


def test_layout_deduplicates_overprinted_text(work_dir: Path):
    from docx import Document
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    source = work_dir / "duplicate.pdf"
    pdf = canvas.Canvas(str(source), pagesize=A4)
    pdf.setFont("Helvetica", 14)
    pdf.drawString(72, 700, "Duplicate")
    pdf.drawString(72, 700, "Duplicate")
    pdf.drawString(72, 670, "Bold")
    pdf.drawString(73.5, 670, "Bold")
    pdf.drawString(72, 640, "Normal repeated ll 0011")
    pdf.showPage()
    pdf.save()
    target = work_dir / "layout-dedup.docx"
    pdf_to_docx_layout(source, target)
    document = Document(str(target))
    texts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    assert sum("Duplicate" in text for text in texts) == 1
    assert all("DuplicateDuplicate" not in text for text in texts)
    assert sum("Bold" in text for text in texts) == 1
    assert any("Normal repeated ll 0011" in text for text in texts)
