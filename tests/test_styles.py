from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

from app.services.styles_service import apply_style, import_style_from_sample


def test_apply_docx_style(docx_file: Path, work_dir: Path):
    target = work_dir / "styled.docx"
    target.write_bytes(docx_file.read_bytes())
    style = {
        "docx": {
            "font_cn": "宋体",
            "font_en": "Times New Roman",
            "body_size": 14,
            "line_spacing": 2.0,
            "margin_left_cm": 2.0,
            "header_text": "测试页眉",
            "footer_text": "测试页脚",
            "page_number": True,
        }
    }
    style["docx"].update({"primary_color": "6D28D9", "accent_color": "8B5CF6", "light_color": "F5F3FF", "heading_color": "4C1D95", "text_color": "2E1065", "header_fill": "6D28D9", "page_background": "F5F3FF"})
    warnings = apply_style(target, "docx", style)
    assert isinstance(warnings, list)
    document = Document(str(target))
    assert round(document.sections[0].left_margin.cm, 1) == 2.0
    assert document.sections[0].header.paragraphs[0].text == "测试页眉"


def test_apply_xlsx_style(xlsx_file: Path, work_dir: Path):
    target = work_dir / "styled.xlsx"
    target.write_bytes(xlsx_file.read_bytes())
    style = {"xlsx": {"font_name": "微软雅黑", "header_fill": "1E3A8A", "header_color": "FFFFFF", "freeze_header": True}}
    style["xlsx"].update({"tab_color": "1E3A8A"})
    apply_style(target, "xlsx", style)
    workbook = load_workbook(target)
    worksheet = workbook.active
    assert worksheet.freeze_panes == "A2"
    assert worksheet["A1"].fill.fgColor.rgb.endswith("1E3A8A")
    assert str(worksheet.sheet_properties.tabColor.rgb).endswith("1E3A8A")
    workbook.close()


def test_apply_pptx_style(pptx_file: Path, work_dir: Path):
    target = work_dir / "styled.pptx"
    target.write_bytes(pptx_file.read_bytes())
    style = {"pptx": {"title_font": "微软雅黑", "body_font": "微软雅黑", "slide_width_inches": 10, "slide_height_inches": 5.625}}
    style["pptx"].update({"background_color": "F5F3FF", "primary_color": "6D28D9", "accent_color": "8B5CF6", "text_color": "2E1065"})
    apply_style(target, "pptx", style)
    presentation = Presentation(str(target))
    assert round(presentation.slide_width / 914400, 2) == 10
    assert abs(presentation.slide_height / 914400 - 5.625) < 0.02
    assert str(presentation.slides[0].background.fill.fore_color.rgb).upper() == "F5F3FF"


def test_import_style_from_docx(docx_file: Path):
    data, warnings = import_style_from_sample(docx_file, "docx", "导入样式")
    assert "docx" in data["applies_to"]
    assert data["docx"]["body_size"] > 0
    assert isinstance(warnings, list)
