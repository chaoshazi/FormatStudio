from pathlib import Path

from docx import Document
from openpyxl import load_workbook

from app.services.table_conv import csv_to_xlsx, xlsx_to_csv
from app.services.text_conv import docx_to_md, md_to_docx, txt_to_md


def test_txt_to_md(txt_file: Path, work_dir: Path):
    target = work_dir / "out.md"
    txt_to_md(txt_file, target)
    assert "第一段内容" in target.read_text(encoding="utf-8")


def test_md_to_docx(md_file: Path, work_dir: Path):
    target = work_dir / "out.docx"
    md_to_docx(md_file, target)
    document = Document(str(target))
    texts = [paragraph.text for paragraph in document.paragraphs]
    assert "标题" in texts
    assert len(document.tables) == 1


def test_docx_to_md(docx_file: Path, work_dir: Path):
    target = work_dir / "out.md"
    docx_to_md(docx_file, target)
    content = target.read_text(encoding="utf-8")
    assert "# 示例文档" in content
    assert "| 项目 | 数量 |" in content


def test_csv_xlsx_roundtrip(csv_file: Path, work_dir: Path):
    xlsx = work_dir / "out.xlsx"
    csv_to_xlsx(csv_file, xlsx, all_text=True)
    workbook = load_workbook(xlsx)
    worksheet = workbook.active
    assert worksheet["A1"].value == "姓名"
    assert str(worksheet["C2"].value) == "001"
    workbook.close()

    csv_out = work_dir / "back.csv"
    xlsx_to_csv(xlsx, csv_out)
    content = csv_out.read_text(encoding="utf-8-sig")
    assert "张三" in content
    assert "001" in content