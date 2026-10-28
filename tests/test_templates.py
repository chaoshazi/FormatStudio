from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

from app.services.template_service import (
    fill_docx,
    fill_pptx,
    fill_xlsx,
    generate_from_template,
    scan_template,
)


def make_docx_template(path: Path) -> Path:
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("姓名：{{na")
    paragraph.add_run("me}}")
    document.add_paragraph("部门：{{dept}}")
    document.add_paragraph("编号：{{code}}")
    document.save(str(path))
    return path


def test_scan_cross_run_placeholder(work_dir: Path):
    template = make_docx_template(work_dir / "template.docx")
    fields = scan_template(template, "docx")
    names = {field["name"] for field in fields}
    assert names == {"name", "dept", "code"}


def test_fill_docx_cross_run(work_dir: Path):
    template = make_docx_template(work_dir / "template.docx")
    missing = fill_docx(template, {"name": "张三", "dept": "研发", "code": "A001"})
    assert missing == set()
    document = Document(str(template))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "姓名：张三" in text
    assert "部门：研发" in text


def test_generate_batch(work_dir: Path):
    template = make_docx_template(work_dir / "template.docx")
    data = work_dir / "data.csv"
    data.write_text("name,dept,code\n张三,研发,A001\n李四,市场,B002\n", encoding="utf-8-sig")
    output = work_dir / "out"
    result = generate_from_template(template, "docx", data, output)
    assert result["generated"] == 2
    assert len(result["outputs"]) == 2
    assert result["archive"] is not None
    names = {path.name for path in result["outputs"]}
    assert "template_001.docx" in names


def test_xlsx_fill(work_dir: Path):
    import openpyxl

    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet["A1"] = "名称"
    worksheet["B1"] = "{{name}}"
    worksheet["A2"] = "部门"
    worksheet["B2"] = "{{dept}}"
    path = work_dir / "template.xlsx"
    workbook.save(path)
    fields = scan_template(path, "xlsx")
    assert {field["name"] for field in fields} == {"name", "dept"}
    missing = fill_xlsx(path, {"name": "张三", "dept": "研发"})
    assert missing == set()
    loaded = load_workbook(path)
    assert loaded.active["B1"].value == "张三"
    loaded.close()


def test_pptx_fill(work_dir: Path):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "{{title}}"
    slide.placeholders[1].text = "负责人：{{owner}}"
    path = work_dir / "template.pptx"
    presentation.save(path)
    fields = scan_template(path, "pptx")
    assert {field["name"] for field in fields} == {"title", "owner"}
    missing = fill_pptx(path, {"title": "周报", "owner": "张三"})
    assert missing == set()
    loaded = Presentation(str(path))
    assert loaded.slides[0].shapes.title.text == "周报"


def test_missing_field_reports(work_dir: Path):
    template = make_docx_template(work_dir / "template.docx")
    data = work_dir / "data.csv"
    data.write_text("name,dept\n张三,研发\n", encoding="utf-8-sig")
    from app.services.template_service import TemplateError

    try:
        generate_from_template(template, "docx", data, work_dir / "out")
    except TemplateError as exc:
        assert exc.code == "MISSING_FIELD"
    else:
        raise AssertionError("expected TemplateError")