from __future__ import annotations

import statistics
import tempfile
from collections import Counter
from pathlib import Path
import pdfplumber
from docx import Document
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image
from pptx import Presentation
from pptx.util import Pt as PptPt

from ..utils import write_text
from .pdf_render import open_pdf, render_page_image
from .table_conv import _docx_add_table


class PdfConversionError(RuntimeError):
    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code
        self.message = message


def _check_password(path: Path, password: str | None) -> None:
    try:
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        if not reader.is_encrypted:
            return
        if not password:
            raise PdfConversionError("PDF_PASSWORD_REQUIRED", "PDF \u5df2\u52a0\u5bc6\uff0c\u8bf7\u63d0\u4f9b\u6253\u5f00\u5bc6\u7801")
        try:
            result = reader.decrypt(password)
        except Exception as exc:
            raise PdfConversionError("PDF_PASSWORD_INVALID", "PDF \u5bc6\u7801\u4e0d\u6b63\u786e") from exc
        if result == 0:
            raise PdfConversionError("PDF_PASSWORD_INVALID", "PDF \u5bc6\u7801\u4e0d\u6b63\u786e")
    except PdfConversionError:
        raise
    except Exception:
        return


def _open_pdf(path: Path, password: str | None = None):
    _check_password(path, password)
    kwargs = {}
    if password:
        kwargs["password"] = password
    try:
        return pdfplumber.open(path, **kwargs)
    except Exception as exc:
        message = str(exc)
        if password or "password" in message.lower() or "encrypted" in message.lower():
            code = "PDF_PASSWORD_INVALID" if password else "PDF_PASSWORD_REQUIRED"
            raise PdfConversionError(code, "PDF \u5df2\u52a0\u5bc6\uff0c\u8bf7\u63d0\u4f9b\u6b63\u786e\u5bc6\u7801") from exc
        raise


def _deduped_chars(page) -> list[dict]:
    """Drop overprinted duplicate characters commonly found in PDFs.

    Some generators draw the same glyph twice at nearly identical coordinates
    to fake bold text or because of an export bug.  extract_words() merges
    those layers and returns duplicated text ("??"), which then becomes
    duplicated Word text boxes.  Deduplicate raw characters by comparing both
    text and bounding-box overlap.
    """
    tolerance = 3.0
    kept: list[dict] = []
    buckets: dict[tuple[str, int, int], list[dict]] = {}

    def overlaps(first: dict, second: dict) -> bool:
        ax0, ax1 = float(first.get("x0", 0)), float(first.get("x1", 0))
        bx0, bx1 = float(second.get("x0", 0)), float(second.get("x1", 0))
        atop, abottom = float(first.get("top", 0)), float(first.get("bottom", 0))
        btop, bbottom = float(second.get("top", 0)), float(second.get("bottom", 0))
        ix = min(ax1, bx1) - max(ax0, bx0)
        iy = min(abottom, bbottom) - max(atop, btop)
        if ix <= 0 or iy <= 0:
            return False
        min_width = min(ax1 - ax0, bx1 - bx0)
        min_height = min(abottom - atop, bbottom - btop)
        return ix >= max(0.8, min_width * 0.5) and iy >= max(0.8, min_height * 0.5)

    for char in page.chars or []:
        text = char.get("text") or ""
        if not text.strip():
            continue
        x0 = float(char.get("x0", 0))
        top = float(char.get("top", 0))
        key = (text, round(x0 / tolerance), round(top / tolerance))
        duplicate = False
        for dx in (-1, 0, 1):
            for dy in (-1, 0, 1):
                for previous in buckets.get((text, key[1] + dx, key[2] + dy), ()):
                    if overlaps(previous, char):
                        duplicate = True
                        break
                if duplicate:
                    break
            if duplicate:
                break
        if duplicate:
            continue
        kept.append(char)
        buckets.setdefault(key, []).append(char)
    return kept


def _chars_to_words(chars: list[dict]) -> list[dict]:
    if not chars:
        return []
    rows: list[list[dict]] = []
    for char in sorted(chars, key=lambda item: (round(float(item.get("top", 0)), 1), item.get("x0", 0))):
        placed = False
        for row in rows:
            if abs(float(row[0].get("top", 0)) - float(char.get("top", 0))) <= 3:
                row.append(char)
                placed = True
                break
        if not placed:
            rows.append([char])
    groups: list[list[dict]] = []
    for row in rows:
        row.sort(key=lambda item: item.get("x0", 0))
        current: list[dict] = []
        for char in row:
            text = char.get("text") or ""
            if text.isspace():
                if current:
                    groups.append(current)
                    current = []
                continue
            if current:
                previous = current[-1]
                gap = float(char.get("x0", 0)) - float(previous.get("x1", 0))
                size = float(char.get("size") or previous.get("size") or 10)
                if gap > max(1.0, size * 0.12):
                    groups.append(current)
                    current = []
            current.append(char)
        if current:
            groups.append(current)
    words: list[dict] = []
    for group in groups:
        text = "".join(char.get("text") or "" for char in group).strip()
        if not text:
            continue
        sizes = [float(char.get("size") or 0) for char in group if char.get("size")]
        fonts = [str(char.get("fontname") or "") for char in group]
        colors = [char.get("non_stroking_color") for char in group if char.get("non_stroking_color")]
        words.append(
            {
                "text": text,
                "x0": min(float(char.get("x0", 0)) for char in group),
                "x1": max(float(char.get("x1", 0)) for char in group),
                "top": min(float(char.get("top", 0)) for char in group),
                "bottom": max(float(char.get("bottom", 0)) for char in group),
                "size": statistics.median(sizes) if sizes else 10.0,
                "font": Counter(fonts).most_common(1)[0][0] if fonts else "",
                "color": colors[0] if colors else None,
            }
        )
    return words


def _page_lines(page) -> list[dict]:
    words = _chars_to_words(_deduped_chars(page))
    if not words:
        return []
    rows: list[list[dict]] = []
    for word in sorted(words, key=lambda item: (round(item["top"], 1), item["x0"])):
        placed = False
        for row in rows:
            if abs(row[0]["top"] - word["top"]) <= 3:
                row.append(word)
                placed = True
                break
        if not placed:
            rows.append([word])
    rows.sort(key=lambda row: min(item["top"] for item in row))
    lines: list[dict] = []
    for row in rows:
        row.sort(key=lambda item: item["x0"])
        text = " ".join(item["text"] for item in row).strip()
        if not text:
            continue
        sizes = [float(item.get("size") or 0) for item in row if item.get("size")]
        fonts = [str(item.get("fontname") or "") for item in row]
        colors = [item.get("non_stroking_color") for item in row if item.get("non_stroking_color")]
        lines.append(
            {
                "text": text,
                "x0": min(item["x0"] for item in row),
                "x1": max(item["x1"] for item in row),
                "top": min(item["top"] for item in row),
                "bottom": max(item["bottom"] for item in row),
                "size": statistics.median(sizes) if sizes else 10.0,
                "font": Counter(fonts).most_common(1)[0][0] if fonts else "",
                "color": colors[0] if colors else None,
            }
        )
    return lines


def _is_bold(font: str) -> bool:
    lowered = font.lower()
    return "bold" in lowered or "black" in lowered or "heavy" in lowered


def _colored_rgb(color) -> RGBColor | None:
    if not color:
        return None
    try:
        values = list(color)
        if len(values) >= 3:
            return RGBColor(
                int(max(0, min(1, values[0])) * 255),
                int(max(0, min(1, values[1])) * 255),
                int(max(0, min(1, values[2])) * 255),
            )
    except Exception:
        return None
    return None


def _paragraphs_from_lines(lines: list[dict]) -> list[dict]:
    if not lines:
        return []
    heights = [max(1.0, line["bottom"] - line["top"]) for line in lines]
    median_height = statistics.median(heights)
    body_size = statistics.median([line["size"] for line in lines]) or 10.0
    blocks: list[dict] = []
    for line in lines:
        if blocks:
            previous = blocks[-1]
            gap = line["top"] - previous["bottom"]
            same_indent = abs(line["x0"] - previous["x0"]) < 24
            if gap <= max(6.0, median_height * 0.9) and same_indent:
                previous["lines"].append(line)
                previous["bottom"] = line["bottom"]
                previous["x1"] = max(previous["x1"], line["x1"])
                continue
        blocks.append(
            {
                "type": "text",
                "lines": [line],
                "x0": line["x0"],
                "x1": line["x1"],
                "top": line["top"],
                "bottom": line["bottom"],
                "size": line["size"],
            }
        )
    for block in blocks:
        block["text"] = "\n".join(line["text"] for line in block["lines"])
        block["sizes"] = [line["size"] for line in block["lines"] if line["size"]]
        block["fonts"] = [line["font"] for line in block["lines"] if line["font"]]
        block["body_size"] = body_size
        block["avg_size"] = statistics.mean(block["sizes"]) if block["sizes"] else body_size
    return blocks


def _inside(bbox: tuple[float, float, float, float], x: float, y: float) -> bool:
    return bbox[0] <= x <= bbox[2] and bbox[1] <= y <= bbox[3]


def _page_tables(page) -> list[dict]:
    results: list[dict] = []
    try:
        tables = page.find_tables()
    except Exception:
        return results
    for table in tables:
        try:
            data = table.extract()
        except Exception:
            continue
        if not data or not any(any(cell for cell in row) for row in data):
            continue
        results.append({"type": "table", "data": data, "bbox": table.bbox})
    return results


def _is_scanned(page) -> bool:
    try:
        return len(page.chars) < 20 and bool(page.images)
    except Exception:
        return False


def _render_crop(page, pdf_path: Path, password: str | None, bbox: tuple[float, float, float, float], out_path: Path) -> Path | None:
    try:
        width = bbox[2] - bbox[0]
        height = bbox[3] - bbox[1]
        if width < 8 or height < 8:
            return None
        pdf = open_pdf(pdf_path, password)
        try:
            pdf_page = pdf[page.page_number - 1]
            bitmap = pdf_page.render(scale=200 / 72)
            image = bitmap.to_pil().convert("RGB")
            scale = 200 / 72
            left = max(0, int(bbox[0] * scale))
            top = max(0, int(bbox[1] * scale))
            right = min(image.width, int(bbox[2] * scale))
            bottom = min(image.height, int(bbox[3] * scale))
            if right - left < 4 or bottom - top < 4:
                return None
            crop = image.crop((left, top, right, bottom))
            out_path.parent.mkdir(parents=True, exist_ok=True)
            crop.save(out_path, "PNG")
            crop.close()
            image.close()
            bitmap.close()
            pdf_page.close()
            return out_path
        finally:
            pdf.close()
    except Exception:
        return None


def _add_picture(doc, image_path: Path, max_width_cm: float = 16.0) -> None:
    try:
        with Image.open(image_path) as image:
            width_px, height_px = image.size
        ratio = height_px / max(1, width_px)
        width_cm = max_width_cm
        if width_cm * ratio > 20:
            width_cm = 20 / ratio
        doc.add_picture(str(image_path), width=Cm(width_cm))
    except Exception:
        pass


def pdf_to_txt(source: Path, destination: Path, password: str | None = None) -> None:
    pdf = _open_pdf(source, password)
    try:
        pages = []
        for page in pdf.pages:
            pages.append(page.extract_text(x_tolerance=2, y_tolerance=3) or "")
        write_text(destination, "\f".join(pages).strip() + "\n")
    finally:
        pdf.close()


def pdf_to_md(source: Path, destination: Path, password: str | None = None) -> None:
    pdf = _open_pdf(source, password)
    try:
        parts = []
        for index, page in enumerate(pdf.pages, start=1):
            text = page.extract_text(x_tolerance=2, y_tolerance=3) or ""
            parts.append(f"<!-- page: {index} -->\n\n{text.strip()}")
        write_text(destination, "\n\n".join(parts).strip() + "\n")
    finally:
        pdf.close()


def pdf_to_xlsx(source: Path, destination: Path, password: str | None = None) -> int:
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    pdf = _open_pdf(source, password)
    count = 0
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    try:
        for page_index, page in enumerate(pdf.pages, start=1):
            for table_index, table in enumerate(_page_tables(page), start=1):
                data = table["data"]
                if not data:
                    continue
                worksheet = workbook.create_sheet(f"P{page_index}-T{table_index}")
                for row_index, row in enumerate(data, start=1):
                    for col_index, value in enumerate(row, start=1):
                        worksheet.cell(row=row_index, column=col_index, value=value)
                        cell = worksheet.cell(row=row_index, column=col_index)
                        cell.alignment = openpyxl.styles.Alignment(vertical="center", wrap_text=True)
                for col in range(1, max(len(row) for row in data) + 1):
                    cell = worksheet.cell(row=1, column=col)
                    cell.font = Font(bold=True, color="1E3A8A")
                    cell.fill = PatternFill("solid", fgColor="DBEAFE")
                    worksheet.column_dimensions[get_column_letter(col)].width = 18
                worksheet.freeze_panes = "A2"
                count += 1
    finally:
        pdf.close()
    if count == 0:
        raise PdfConversionError("NO_TABLE", "未在 PDF 中识别到表格；请确认 PDF 含文本层或改用其他格式")
    workbook.save(str(destination))
    return count


def pdf_to_docx_editable(source: Path, destination: Path, password: str | None = None) -> list[str]:
    pdf = _open_pdf(source, password)
    warnings: list[str] = []
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)
    normal = document.styles["Normal"]
    normal.font.name = "微软雅黑"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    scanned_pages: list[int] = []
    try:
        total = len(pdf.pages)
        for page_index, page in enumerate(pdf.pages, start=1):
            if page_index > 1:
                document.add_page_break()
            if _is_scanned(page):
                scanned_pages.append(page_index)
                target = Path(tempfile.gettempdir()) / f"pdf-scan-{source.stem}-{page_index}.png"
                render_page_image(source, target, page_index - 1, 180, password)
                _add_picture(document, target, 16)
                continue
            lines = _page_lines(page)
            blocks = _paragraphs_from_lines(lines)
            tables = _page_tables(page)
            for table in tables:
                bbox = table["bbox"]
                for block in blocks:
                    filtered = [
                        line
                        for line in block["lines"]
                        if not _inside(bbox, (line["x0"] + line["x1"]) / 2, (line["top"] + line["bottom"]) / 2)
                    ]
                    if filtered:
                        block["lines"] = filtered
                        block["text"] = "\n".join(line["text"] for line in filtered)
                blocks.append(
                    {
                        "type": "table",
                        "data": table["data"],
                        "top": bbox[1],
                        "x0": bbox[0],
                        "bottom": bbox[3],
                    }
                )
            image_blocks = []
            for image_index, image in enumerate(page.images):
                bbox = (image["x0"], image["top"], image["x1"], image["bottom"])
                if (bbox[2] - bbox[0]) * (bbox[3] - bbox[1]) < 500:
                    continue
                target = Path(tempfile.gettempdir()) / f"pdf-img-{source.stem}-{page_index}-{image_index}.png"
                cropped = _render_crop(page, source, password, bbox, target)
                if cropped:
                    image_blocks.append({"type": "image", "path": cropped, "top": bbox[1], "x0": bbox[0]})
            blocks = [block for block in blocks if block.get("type") != "text" or block.get("text", "").strip()]
            blocks.extend(image_blocks)
            blocks.sort(key=lambda block: (round(block.get("top", 0), 1), block.get("x0", 0)))
            for block in blocks:
                if block["type"] == "table":
                    _docx_add_table(document, block["data"])
                    document.add_paragraph()
                elif block["type"] == "image":
                    _add_picture(document, block["path"])
                else:
                    text = block.get("text", "").strip()
                    if not text:
                        continue
                    avg_size = block.get("avg_size", 11)
                    body_size = block.get("body_size", 11)
                    if avg_size >= body_size * 1.35 and len(text) <= 100:
                        paragraph = document.add_heading(level=1)
                        _add_text_runs(paragraph, block, Pt(18), bold=True)
                    elif avg_size >= body_size * 1.15 and len(text) <= 100:
                        paragraph = document.add_heading(level=2)
                        _add_text_runs(paragraph, block, Pt(15), bold=True)
                    else:
                        paragraph = document.add_paragraph()
                        paragraph.paragraph_format.line_spacing = 1.5
                        paragraph.paragraph_format.space_after = Pt(6)
                        _add_text_runs(paragraph, block, Pt(max(8, min(14, avg_size))), bold=False)
        if scanned_pages and len(scanned_pages) == total:
            raise PdfConversionError("NEEDS_OCR", "整份 PDF 均为扫描页，没有可提取的文本层，需要 OCR 后再转换")
        if scanned_pages:
            warnings.append(
                f"第 {', '.join(map(str, scanned_pages))} 页为扫描页，已按图片插入；可编辑文本仅来自其余页面"
            )
    finally:
        pdf.close()
    document.save(str(destination))
    return warnings


def _add_text_runs(paragraph, block: dict, size: Pt, bold: bool) -> None:
    for line_index, line in enumerate(block["lines"]):
        run = paragraph.add_run(line["text"])
        run.font.name = _font_name(line.get("font", ""))
        run.font.size = size
        run.bold = bold or _is_bold(line.get("font", ""))
        rgb = _colored_rgb(line.get("color"))
        if rgb and not (rgb[0] == 0 and rgb[1] == 0 and rgb[2] == 0):
            run.font.color.rgb = rgb
        rpr = run._element.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = rpr.makeelement(qn("w:rFonts"), {})
            rpr.append(fonts)
        fonts.set(qn("w:eastAsia"), run.font.name or "微软雅黑")
        if line_index < len(block["lines"]) - 1:
            run.add_break()


def _font_name(font: str) -> str:
    cleaned = font.split("+")[-1] if font else ""
    cleaned = cleaned.split("-")[0] if cleaned else ""
    if not cleaned:
        return "微软雅黑"
    lowered = cleaned.lower()
    if "song" in lowered or lowered in {"simsun", "nsimsun"}:
        return "宋体"
    if "hei" in lowered:
        return "黑体"
    if "kai" in lowered:
        return "楷体"
    if "fang" in lowered:
        return "仿宋"
    if "yahei" in lowered or "msyh" in lowered:
        return "微软雅黑"
    return cleaned


def _mask_text(image: Image.Image, lines: list[dict], scale: float) -> Image.Image:
    from PIL import ImageDraw

    draw = ImageDraw.Draw(image)
    for line in lines:
        left = max(0, int(line["x0"] * scale) - 2)
        top = max(0, int(line["top"] * scale) - 2)
        right = min(image.width, int(line["x1"] * scale) + 2)
        bottom = min(image.height, int(line["bottom"] * scale) + 2)
        if right <= left or bottom <= top:
            continue
        samples = []
        for x in range(left, right, max(1, (right - left) // 8)):
            for y in (max(0, top - 3), min(image.height - 1, bottom + 3)):
                samples.append(image.getpixel((x, y)))
        for y in range(top, bottom, max(1, (bottom - top) // 4)):
            for x in (max(0, left - 3), min(image.width - 1, right + 3)):
                samples.append(image.getpixel((x, y)))
        if samples:
            color = tuple(int(statistics.median(channel)) for channel in zip(*samples))
        else:
            color = (255, 255, 255)
        draw.rectangle([left, top, right, bottom], fill=color)
    return image


def _set_frame(paragraph, x_twips: int, y_twips: int, w_twips: int, h_twips: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    frame = p_pr.find(qn("w:framePr"))
    if frame is None:
        frame = OxmlElement("w:framePr")
        p_pr.insert(0, frame)
    frame.set(qn("w:w"), str(max(1, w_twips)))
    frame.set(qn("w:h"), str(max(1, h_twips)))
    frame.set(qn("w:x"), str(max(0, x_twips)))
    frame.set(qn("w:y"), str(max(0, y_twips)))
    frame.set(qn("w:hAnchor"), "page")
    frame.set(qn("w:vAnchor"), "page")
    frame.set(qn("w:wrap"), "none")
    frame.set(qn("w:dropCap"), "none")
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0


def pdf_to_docx_layout(source: Path, destination: Path, password: str | None = None) -> list[str]:
    pdf = _open_pdf(source, password)
    document = Document()
    warnings: list[str] = []
    total = len(pdf.pages)
    scanned_pages: list[int] = []
    try:
        for page_index, page in enumerate(pdf.pages):
            page_number = page_index + 1
            lines = [] if _is_scanned(page) else _page_lines(page)
            if _is_scanned(page):
                scanned_pages.append(page_number)
            width = float(page.width)
            height = float(page.height)
            if page_index == 0:
                section = document.sections[0]
            else:
                section = document.add_section(WD_SECTION.NEW_PAGE)
            section.page_width = Pt(width)
            section.page_height = Pt(height)
            section.top_margin = Pt(0)
            section.bottom_margin = Pt(0)
            section.left_margin = Pt(0)
            section.right_margin = Pt(0)
            section.header_distance = Pt(0)
            section.footer_distance = Pt(0)
            section.header.is_linked_to_previous = False
            section.footer.is_linked_to_previous = False
            target = Path(tempfile.gettempdir()) / f"pdf-layout-{source.stem}-{page_number}.png"
            render_page_image(source, target, page_index, 200, password)
            with Image.open(target) as image:
                image = image.convert("RGB")
                scale = image.width / max(1.0, width)
                if lines:
                    image = _mask_text(image, lines, scale)
                image.save(target, "PNG")
                image.close()
            header = section.header
            header_paragraph = header.paragraphs[0]
            header_paragraph.paragraph_format.space_before = Pt(0)
            header_paragraph.paragraph_format.space_after = Pt(0)
            header_paragraph.paragraph_format.line_spacing = 1.0
            run = header_paragraph.add_run()
            run.add_picture(str(target), width=Pt(width), height=Pt(height))
            if not lines:
                continue
            for line in lines:
                paragraph = document.add_paragraph()
                x = int(line["x0"] * 20)
                y = int(line["top"] * 20)
                w = int(max(line["x1"] - line["x0"] + 4, width - line["x0"] - 4) * 20)
                h = int((line["bottom"] - line["top"]) * 20) + 80
                _set_frame(paragraph, x, y, w, h)
                for run_index, part in enumerate(line["text"].split("\n")):
                    run = paragraph.add_run(part)
                    run.font.size = Pt(max(1, line["size"]))
                    run.font.name = _font_name(line.get("font", ""))
                    run.bold = _is_bold(line.get("font", ""))
                    rgb = _colored_rgb(line.get("color"))
                    if rgb:
                        run.font.color.rgb = rgb
        if scanned_pages:
            warnings.append(
                f"第 {', '.join(map(str, scanned_pages))} 页没有文本层，已保留为页面图片，无法直接编辑"
            )
    finally:
        pdf.close()
    if total == 0:
        raise PdfConversionError("EMPTY_PDF", "PDF 没有页面")
    document.save(str(destination))
    return warnings


def pdf_to_pptx_editable(source: Path, destination: Path, password: str | None = None) -> list[str]:
    pdf = _open_pdf(source, password)
    presentation = Presentation()
    if len(pdf.pages) > 0:
        first = pdf.pages[0]
        presentation.slide_width = PptPt(float(first.width))
        presentation.slide_height = PptPt(float(first.height))
    layout = presentation.slide_layouts[6]
    warnings: list[str] = []
    scanned_pages: list[int] = []
    try:
        total = len(pdf.pages)
        for page_index, page in enumerate(pdf.pages, start=1):
            slide = presentation.slides.add_slide(layout)
            if _is_scanned(page):
                scanned_pages.append(page_index)
                target = Path(tempfile.gettempdir()) / f"pdf-ppt-{source.stem}-{page_index}.png"
                render_page_image(source, target, page_index - 1, 180, password)
                slide.shapes.add_picture(str(target), 0, 0, width=presentation.slide_width, height=presentation.slide_height)
                continue
            lines = _page_lines(page)
            blocks = _paragraphs_from_lines(lines)
            tables = _page_tables(page)
            for table in tables:
                bbox = table["bbox"]
                for block in blocks:
                    filtered = [
                        line
                        for line in block["lines"]
                        if not _inside(bbox, (line["x0"] + line["x1"]) / 2, (line["top"] + line["bottom"]) / 2)
                    ]
                    if filtered:
                        block["lines"] = filtered
                        block["text"] = "\n".join(line["text"] for line in filtered)
                blocks.append({"type": "table", "data": table["data"], "top": bbox[1], "x0": bbox[0], "bbox": bbox})
            image_blocks = []
            for image_index, image in enumerate(page.images):
                bbox = (image["x0"], image["top"], image["x1"], image["bottom"])
                if (bbox[2] - bbox[0]) * (bbox[3] - bbox[1]) < 500:
                    continue
                target = Path(tempfile.gettempdir()) / f"pdf-ppt-img-{source.stem}-{page_index}-{image_index}.png"
                cropped = _render_crop(page, source, password, bbox, target)
                if cropped:
                    image_blocks.append({"type": "image", "path": cropped, "top": bbox[1], "x0": bbox[0], "bbox": bbox})
            blocks = [block for block in blocks if block.get("type") != "text" or block.get("text", "").strip()]
            blocks.extend(image_blocks)
            blocks.sort(key=lambda block: (round(block.get("top", 0), 1), block.get("x0", 0)))
            for block in blocks:
                if block["type"] == "table":
                    data = block["data"]
                    if not data:
                        continue
                    rows = len(data)
                    cols = max(len(row) for row in data)
                    bbox = block["bbox"]
                    shape = slide.shapes.add_table(
                        rows,
                        cols,
                        PptPt(bbox[0]),
                        PptPt(bbox[1]),
                        PptPt(bbox[2] - bbox[0]),
                        PptPt(bbox[3] - bbox[1]),
                    )
                    for row_index, row in enumerate(data):
                        for col_index in range(cols):
                            value = row[col_index] if col_index < len(row) else ""
                            cell = shape.table.cell(row_index, col_index)
                            cell.text = str(value or "")
                            for paragraph in cell.text_frame.paragraphs:
                                for run in paragraph.runs:
                                    run.font.size = PptPt(8)
                elif block["type"] == "image":
                    bbox = block["bbox"]
                    try:
                        slide.shapes.add_picture(
                            str(block["path"]),
                            PptPt(bbox[0]),
                            PptPt(bbox[1]),
                            width=PptPt(bbox[2] - bbox[0]),
                        )
                    except Exception:
                        pass
                else:
                    text = block.get("text", "").strip()
                    if not text:
                        continue
                    width = max(20, block["x1"] - block["x0"] + 4)
                    height = max(12, block["bottom"] - block["top"] + 4)
                    shape = slide.shapes.add_textbox(
                        PptPt(block["x0"]),
                        PptPt(block["top"]),
                        PptPt(width),
                        PptPt(height),
                    )
                    text_frame = shape.text_frame
                    text_frame.word_wrap = False
                    text_frame.margin_left = 0
                    text_frame.margin_right = 0
                    text_frame.margin_top = 0
                    text_frame.margin_bottom = 0
                    paragraph = text_frame.paragraphs[0]
                    run = paragraph.add_run()
                    run.text = text
                    run.font.size = PptPt(max(4, block.get("avg_size", 11)))
                    run.font.name = _font_name(block.get("fonts", [""])[0] if block.get("fonts") else "")
                    run.font.bold = _is_bold(block.get("fonts", [""])[0] if block.get("fonts") else "")
        if scanned_pages and len(scanned_pages) == total:
            warnings.append("整份 PDF 均为扫描页，PPT 为图片版，需 OCR 后才能提取文字")
        elif scanned_pages:
            warnings.append(f"第 {', '.join(map(str, scanned_pages))} 页为扫描页，已按整页图片放入幻灯片")
    finally:
        pdf.close()
    presentation.save(str(destination))
    return warnings


def pdf_to_pptx_layout(source: Path, destination: Path, password: str | None = None) -> list[str]:
    pdf = _open_pdf(source, password)
    presentation = Presentation()
    if len(pdf.pages) > 0:
        first = pdf.pages[0]
        presentation.slide_width = PptPt(float(first.width))
        presentation.slide_height = PptPt(float(first.height))
    layout = presentation.slide_layouts[6]
    try:
        for page_index in range(len(pdf.pages)):
            slide = presentation.slides.add_slide(layout)
            target = Path(tempfile.gettempdir()) / f"pdf-ppt-layout-{source.stem}-{page_index + 1}.png"
            render_page_image(source, target, page_index, 200, password)
            slide.shapes.add_picture(
                str(target),
                0,
                0,
                width=presentation.slide_width,
                height=presentation.slide_height,
            )
    finally:
        pdf.close()
    presentation.save(str(destination))
    return []
