from __future__ import annotations

from pathlib import Path
from typing import Any

from ..utils import detect_and_read_text


def _safe(callable_, fallback: Any = 0) -> Any:
    try:
        return callable_()
    except Exception:
        return fallback


def inspect_document(path: Path, fmt: str) -> dict[str, Any]:
    meta: dict[str, Any] = {"size": path.stat().st_size if path.exists() else 0}
    if not path.exists():
        return meta
    if fmt == "pdf":
        try:
            import pypdfium2 as pdfium

            pdf = pdfium.PdfDocument(str(path))
            meta["pages"] = len(pdf)
            pdf.close()
        except Exception:
            meta["pages"] = 0
    elif fmt == "docx":
        try:
            import docx

            document = docx.Document(str(path))
            paragraphs = [p for p in document.paragraphs if p.text.strip()]
            meta["paragraphs"] = len(paragraphs)
            meta["tables"] = len(document.tables)
            meta["words"] = sum(len(p.text.split()) for p in paragraphs)
        except Exception:
            pass
    elif fmt == "xlsx":
        try:
            import openpyxl

            workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
            meta["sheets"] = len(workbook.worksheets)
            meta["rows"] = sum(worksheet.max_row for worksheet in workbook.worksheets)
            workbook.close()
        except Exception:
            pass
    elif fmt == "pptx":
        try:
            from pptx import Presentation

            presentation = Presentation(str(path))
            meta["slides"] = len(presentation.slides)
        except Exception:
            pass
    elif fmt in {"txt", "md", "csv"}:
        try:
            text, _ = detect_and_read_text(path)
            meta["lines"] = text.count("\n") + (1 if text else 0)
            meta["chars"] = len(text)
        except Exception:
            pass
    return meta


def merge_meta(*parts: dict[str, Any]) -> dict[str, Any]:
    result: dict[str, Any] = {}
    for part in parts:
        result.update({key: value for key, value in part.items() if value is not None})
    return result