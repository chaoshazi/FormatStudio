from __future__ import annotations

from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor as PptRGB
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt as PptPt


def build_pptx_rich(spec: dict[str, Any], path: Path, palettes: dict[str, dict[str, str]], deck_index: int) -> None:
    palette = palettes[spec["palette"]]
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    blank = presentation.slide_layouts[6]
    width = presentation.slide_width
    height = presentation.slide_height

    def rgb(value: str):
        return PptRGB.from_string(value.lstrip("#"))

    def rect(slide, x, y, w, h, color, shape=MSO_SHAPE.RECTANGLE):
        item = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        item.fill.solid()
        item.fill.fore_color.rgb = rgb(color)
        item.line.fill.background()
        return item

    def outline(slide, x, y, w, h, color, shape=MSO_SHAPE.OVAL):
        item = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        item.fill.background()
        item.line.color.rgb = rgb(color)
        item.line.width = PptPt(2)
        return item

    def text(slide, value, x, y, w, h, size, color, bold=False, align=PP_ALIGN.LEFT, font="微软雅黑", spacing=1.0):
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        frame = box.text_frame
        frame.word_wrap = True
        frame.margin_left = 0
        frame.margin_right = 0
        frame.margin_top = 0
        frame.margin_bottom = 0
        frame.vertical_anchor = MSO_ANCHOR.TOP
        lines = str(value).split("\n")
        for index, line in enumerate(lines):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.alignment = align
            paragraph.line_spacing = spacing
            run = paragraph.add_run()
            run.text = line
            run.font.size = PptPt(size)
            run.font.bold = bold
            run.font.name = font
            run.font.color.rgb = rgb(color)
        return box

    def footer(slide, page: int) -> None:
        rect(slide, 0.9, 6.98, 11.55, 0.01, palette["light"])
        text(slide, spec["subtitle"], 0.9, 7.08, 8.5, 0.22, 8, palette["muted"])
        text(slide, f"{page:02d}", 12.35, 7.08, 0.4, 0.22, 8, palette["muted"], align=PP_ALIGN.RIGHT)

    def title_dark() -> None:
        slide = presentation.slides.add_slide(blank)
        rect(slide, 0, 0, 13.333, 7.5, palette["dark"])
        rect(slide, 0, 0, 0.24, 7.5, palette["accent"])
        rect(slide, 10.7, -0.8, 3.2, 3.2, palette["primary"], MSO_SHAPE.OVAL)
        outline(slide, 9.65, 4.9, 2.5, 2.5, palette["accent"])
        text(slide, "FORMAT STUDIO  /  {{company}}", 0.9, 0.85, 8.5, 0.3, 9, palette["accent"], True)
        text(slide, spec["title"], 0.9, 2.25, 9.6, 1.55, 36, palette["white"], True, spacing=1.05)
        text(slide, spec["subtitle"], 0.95, 3.95, 8.8, 0.5, 16, palette["accent"])
        rect(slide, 0.95, 4.75, 2.4, 0.06, palette["accent"])
        text(slide, "{{presenter}}  |  {{date}}", 0.95, 6.4, 7.0, 0.3, 10, palette["muted"])

    def title_split() -> None:
        slide = presentation.slides.add_slide(blank)
        rect(slide, 0, 0, 13.333, 7.5, palette["light"])
        rect(slide, 0, 0, 6.1, 7.5, palette["dark"])
        rect(slide, 6.1, 0, 0.16, 7.5, palette["accent"])
        rect(slide, 10.1, 5.4, 2.5, 2.5, palette["white"], MSO_SHAPE.OVAL)
        text(slide, "FORMAT STUDIO", 0.8, 0.85, 4.5, 0.3, 9, palette["accent"], True)
        text(slide, spec["name"].split(" · ")[0], 0.8, 1.65, 4.6, 1.2, 25, palette["white"], True)
        text(slide, "{{company}}\n{{date}}", 0.82, 5.7, 4.2, 0.8, 10, palette["light"], spacing=1.3)
        text(slide, "PRESENTATION", 6.7, 1.0, 5.4, 0.3, 9, palette["primary"], True)
        text(slide, spec["title"], 6.7, 1.75, 5.6, 2.1, 30, palette["dark"], True, spacing=1.05)
        text(slide, spec["subtitle"], 6.72, 4.0, 5.3, 0.6, 13, palette["muted"])
        rect(slide, 6.72, 4.9, 1.8, 0.05, palette["primary"])
        text(slide, "{{presenter}}", 6.72, 5.2, 4.5, 0.3, 10, palette["dark"])

    def title_light() -> None:
        slide = presentation.slides.add_slide(blank)
        rect(slide, 0, 0, 13.333, 7.5, palette["white"])
        rect(slide, 0, 0, 13.333, 0.22, palette["primary"])
        rect(slide, 0.9, 1.05, 1.1, 0.08, palette["accent"])
        text(slide, spec["name"].split(" · ")[0], 0.9, 1.3, 6.0, 0.35, 10, palette["primary"], True)
        text(slide, spec["title"], 0.9, 2.0, 9.4, 1.8, 38, palette["dark"], True, spacing=1.05)
        text(slide, spec["subtitle"], 0.92, 4.15, 8.6, 0.55, 14, palette["muted"])
        for index, color in enumerate((palette["primary"], palette["accent"], palette["light"])):
            rect(slide, 0.92 + index * 1.25, 5.15, 1.0, 0.12, color)
        text(slide, "{{presenter}}  |  {{date}}  |  {{company}}", 0.92, 6.35, 8.0, 0.3, 10, palette["muted"])

    variants = [title_dark, title_split, title_light]
    variants[deck_index % len(variants)]()
    def slide_header(slide, slide_spec, page: int) -> None:
        rect(slide, 0, 0, 13.333, 0.1, palette["primary"])
        text(slide, f"{page - 1:02d}  /  {spec['name'].split(' · ')[0]}", 0.9, 0.48, 6.0, 0.28, 8.5, palette["muted"], True)
        text(slide, slide_spec["title"], 0.9, 0.85, 10.8, 0.8, 27, palette["dark"], True)
        rect(slide, 0.92, 1.66, 1.5, 0.05, palette["accent"])

    def bullets_layout(slide, slide_spec, page: int) -> None:
        rect(slide, 0, 0, 13.333, 7.5, palette["white"])
        slide_header(slide, slide_spec, page)
        bullets = slide_spec.get("bullets", [])
        for index, bullet in enumerate(bullets[:4]):
            y = 2.15 + index * 1.02
            rect(slide, 0.95, y + 0.05, 0.5, 0.5, palette["primary"], MSO_SHAPE.OVAL)
            text(slide, f"{index + 1:02d}", 0.95, y + 0.13, 0.5, 0.25, 8, palette["white"], True, align=PP_ALIGN.CENTER)
            text(slide, bullet, 1.65, y, 6.9, 0.7, 15, palette["dark"])
            rect(slide, 1.65, y + 0.68, 6.5, 0.01, palette["light"])
        rect(slide, 8.7, 2.1, 3.75, 4.35, palette["light"], MSO_SHAPE.ROUNDED_RECTANGLE)
        text(slide, f"{page:02d}", 9.0, 2.4, 3.0, 0.9, 46, palette["primary"], True)
        text(slide, "{{company}}", 9.02, 3.5, 3.0, 0.35, 12, palette["dark"], True)
        text(slide, "{{presenter}}  ·  {{date}}", 9.02, 3.9, 3.0, 0.5, 10, palette["muted"])
        outline(slide, 10.7, 5.0, 1.2, 1.2, palette["accent"])
        footer(slide, page)

    def kpis_layout(slide, slide_spec, page: int) -> None:
        rect(slide, 0, 0, 13.333, 7.5, palette["light"])
        slide_header(slide, slide_spec, page)
        for index, (label, value) in enumerate(slide_spec.get("kpis", [])[:3]):
            x = 0.9 + index * 3.95
            rect(slide, x, 2.0, 3.45, 1.65, palette["white"], MSO_SHAPE.ROUNDED_RECTANGLE)
            rect(slide, x, 2.0, 0.08, 1.65, palette["primary"])
            text(slide, label, x + 0.28, 2.25, 2.8, 0.3, 10, palette["muted"])
            text(slide, value, x + 0.28, 2.62, 2.9, 0.75, 27, palette["primary"], True)
        bullets = slide_spec.get("bullets", [])
        for index, bullet in enumerate(bullets[:3]):
            y = 4.25 + index * 0.6
            rect(slide, 0.95, y + 0.11, 0.12, 0.12, palette["accent"], MSO_SHAPE.OVAL)
            text(slide, bullet, 1.25, y, 6.5, 0.45, 13, palette["dark"])
        bars = [0.45, 0.72, 0.58, 0.86]
        for index, ratio in enumerate(bars):
            x = 8.9 + index * 0.9
            rect(slide, x, 6.0 - ratio * 2.4, 0.52, ratio * 2.4, palette["primary"])
            text(slide, str(index + 1), x, 6.08, 0.5, 0.25, 8, palette["muted"], align=PP_ALIGN.CENTER)
        footer(slide, page)

    def timeline_layout(slide, slide_spec, page: int) -> None:
        rect(slide, 0, 0, 13.333, 7.5, palette["white"])
        slide_header(slide, slide_spec, page)
        bullets = slide_spec.get("bullets", [])
        rect(slide, 1.0, 3.85, 11.2, 0.03, palette["light"])
        for index, bullet in enumerate(bullets[:4]):
            x = 1.2 + index * 3.35
            rect(slide, x, 3.58, 0.58, 0.58, palette["primary"], MSO_SHAPE.OVAL)
            text(slide, f"{index + 1}", x, 3.71, 0.58, 0.25, 10, palette["white"], True, align=PP_ALIGN.CENTER)
            text(slide, f"0{index + 1}", x, 2.55, 1.4, 0.35, 12, palette["accent"], True)
            text(slide, bullet, x, 4.4, 2.9, 0.9, 12.5, palette["dark"])
        rect(slide, 0.9, 6.35, 11.5, 0.01, palette["light"])
        footer(slide, page)

    def split_layout(slide, slide_spec, page: int) -> None:
        rect(slide, 0, 0, 13.333, 7.5, palette["white"])
        rect(slide, 7.7, 0, 5.633, 7.5, palette["dark"])
        rect(slide, 7.7, 0, 0.12, 7.5, palette["accent"])
        slide_header(slide, slide_spec, page)
        bullets = slide_spec.get("bullets", [])
        quote = bullets[0] if bullets else ""
        text(slide, "“", 7.95, 0.85, 1.2, 0.9, 58, palette["accent"], True)
        text(slide, quote, 8.05, 2.1, 4.4, 1.8, 20, palette["white"], True, spacing=1.2)
        for index, bullet in enumerate(bullets[1:4]):
            y = 4.45 + index * 0.62
            rect(slide, 8.1, y + 0.1, 0.1, 0.1, palette["accent"], MSO_SHAPE.OVAL)
            text(slide, bullet, 8.35, y, 4.0, 0.5, 11.5, palette["light"])
        text(slide, "{{company}}  ·  {{presenter}}", 0.95, 6.3, 6.0, 0.3, 9, palette["muted"])
        footer(slide, page)

    def table_layout(slide, slide_spec, page: int) -> None:
        rect(slide, 0, 0, 13.333, 7.5, palette["white"])
        slide_header(slide, slide_spec, page)
        bullets = slide_spec.get("bullets", [])
        rows = len(bullets) + 1
        table = slide.shapes.add_table(rows, 2, Inches(0.95), Inches(2.0), Inches(10.9), Inches(0.55 * rows)).table
        table.columns[0].width = Inches(2.2)
        table.columns[1].width = Inches(8.7)
        table.cell(0, 0).text = "NO."
        table.cell(0, 1).text = "KEY POINT"
        for row in range(rows):
            for col in range(2):
                cell = table.cell(row, col)
                cell.fill.solid()
                cell.fill.fore_color.rgb = rgb(palette["primary"] if row == 0 else (palette["light"] if row % 2 else palette["white"]))
                for paragraph in cell.text_frame.paragraphs:
                    for run in paragraph.runs:
                        run.font.name = "微软雅黑"
                        run.font.size = PptPt(10)
                        run.font.bold = row == 0
                        run.font.color.rgb = rgb(palette["white"] if row == 0 else palette["dark"])
        for index, bullet in enumerate(bullets[:4], start=1):
            table.cell(index, 0).text = f"{index:02d}"
            table.cell(index, 1).text = bullet
        footer(slide, page)

    layouts = [bullets_layout, split_layout, timeline_layout, table_layout]
    for index, slide_spec in enumerate(spec["slides"], start=2):
        slide = presentation.slides.add_slide(blank)
        if slide_spec.get("kpis"):
            kpis_layout(slide, slide_spec, index)
        elif slide_spec.get("layout") == "timeline":
            timeline_layout(slide, slide_spec, index)
        elif slide_spec.get("layout") == "table":
            table_layout(slide, slide_spec, index)
        else:
            layouts[(deck_index + index) % len(layouts)](slide, slide_spec, index)
    closing = presentation.slides.add_slide(blank)
    rect(closing, 0, 0, 13.333, 7.5, palette["dark"])
    rect(closing, 0, 0, 0.24, 7.5, palette["accent"])
    rect(closing, 10.8, 1.0, 2.2, 2.2, palette["primary"], MSO_SHAPE.OVAL)
    outline(closing, 0.9, 5.5, 2.2, 2.2, palette["accent"])
    text(closing, "谢谢", 0.95, 2.5, 6.0, 1.0, 46, palette["white"], True)
    text(closing, "{{company}}  |  {{presenter}}", 0.98, 3.85, 8.0, 0.4, 15, palette["accent"])
    text(closing, "{{contact}}", 0.98, 4.45, 8.0, 0.35, 12, palette["light"])
    path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(str(path))