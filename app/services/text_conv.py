from __future__ import annotations

import html
import re
from html.parser import HTMLParser
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from ..utils import detect_and_read_text, write_text


class _HTMLText(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []

    def handle_data(self, data: str) -> None:
        self.parts.append(data)

    def handle_starttag(self, tag: str, attrs) -> None:
        if tag in {"p", "br", "div", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6"}:
            self.parts.append("\n")

    def text(self) -> str:
        value = "".join(self.parts)
        value = re.sub(r"\n{3,}", "\n\n", value)
        return value.strip() + "\n"


def txt_to_md(source: Path, destination: Path) -> None:
    text, _ = detect_and_read_text(source)
    write_text(destination, text.strip() + "\n")


def txt_to_docx(source: Path, destination: Path) -> None:
    text, _ = detect_and_read_text(source)
    document = Document()
    blocks = re.split(r"\n\s*\n", text.replace("\r\n", "\n").replace("\r", "\n"))
    for block in blocks:
        if not block.strip():
            continue
        paragraph = document.add_paragraph()
        paragraph.add_run(block.strip())
        paragraph.paragraph_format.space_after = Pt(6)
        paragraph.paragraph_format.line_spacing = 1.5
    document.save(str(destination))


def _markdown_plain_text(text: str) -> str:
    from markdown_it import MarkdownIt

    md = MarkdownIt("commonmark").enable(["table", "strikethrough"])
    parser = _HTMLText()
    parser.feed(md.render(text))
    value = html.unescape(parser.text())
    value = re.sub(r"[ \t]+\n", "\n", value)
    value = re.sub(r"\n{3,}", "\n\n", value)
    return value


def md_to_txt(source: Path, destination: Path) -> None:
    text, _ = detect_and_read_text(source)
    write_text(destination, _markdown_plain_text(text))
def _add_inline(paragraph, tokens, base_dir: Path) -> None:
    bold = False
    italic = False
    link_stack: list[str] = []
    for token in tokens:
        kind = token.type
        if kind == "strong_open":
            bold = True
        elif kind == "strong_close":
            bold = False
        elif kind == "em_open":
            italic = True
        elif kind == "em_close":
            italic = False
        elif kind == "code_inline":
            run = paragraph.add_run(token.content)
            run.font.name = "Consolas"
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(0xB9, 0x1C, 0x1C)
        elif kind == "link_open":
            link_stack.append(token.attrGet("href") or "")
        elif kind == "link_close":
            if link_stack:
                href = link_stack.pop()
                if href and not href.startswith("#"):
                    run = paragraph.add_run(f" ({href})")
                    run.font.size = Pt(9)
                    run.font.color.rgb = RGBColor(0x60, 0x6B, 0x7A)
        elif kind == "image":
            src = token.attrGet("src") or ""
            alt = token.content or "图片"
            target = None
            if src and not re.match(r"^[a-z]+://", src, re.I):
                candidate = (base_dir / src).resolve()
                if candidate.is_file():
                    target = candidate
            if target is not None:
                try:
                    paragraph.add_run().add_picture(str(target), width=Cm(12))
                    continue
                except Exception:
                    pass
            paragraph.add_run(f"[图片: {alt}]")
        elif kind in {"softbreak", "hardbreak"}:
            paragraph.add_run().add_break()
        elif kind in {"text", "html_inline"}:
            if not token.content:
                continue
            run = paragraph.add_run(token.content)
            run.bold = bold
            run.italic = italic
            if link_stack:
                run.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)
                run.underline = True


def md_to_docx(source: Path, destination: Path) -> None:
    from markdown_it import MarkdownIt

    text, _ = detect_and_read_text(source)
    md = MarkdownIt("commonmark").enable(["table", "strikethrough"])
    tokens = md.parse(text)
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "微软雅黑"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

    base_dir = source.parent
    list_stack: list[dict] = []
    table_rows: list[list] = []
    current_row: list | None = None
    current_cell: list | None = None
    quote_depth = 0
    heading_level = 0
    pending_prefix: tuple[str, int] | None = None

    for index, token in enumerate(tokens):
        kind = token.type
        if kind == "heading_open":
            heading_level = int(token.tag[1]) if token.tag[1:].isdigit() else 1
            paragraph = document.add_heading(level=min(heading_level, 6))
            children = tokens[index + 1].children if index + 1 < len(tokens) else []
            _add_inline(paragraph, children or [], base_dir)
        elif kind == "heading_close":
            heading_level = 0
        elif kind == "paragraph_open":
            paragraph = document.add_paragraph()
            if pending_prefix is not None:
                prefix, depth = pending_prefix
                paragraph.paragraph_format.left_indent = Cm(0.74 * depth)
                paragraph.paragraph_format.first_line_indent = Cm(-0.4)
                paragraph.add_run(prefix)
                pending_prefix = None
            paragraph.paragraph_format.space_after = Pt(6)
            paragraph.paragraph_format.line_spacing = 1.5
            if quote_depth:
                paragraph.paragraph_format.left_indent = Cm(0.8 * quote_depth)
            children = tokens[index + 1].children if index + 1 < len(tokens) else []
            _add_inline(paragraph, children or [], base_dir)
            if quote_depth:
                for run in paragraph.runs:
                    run.italic = True
        elif kind in {"fence", "code_block"}:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(4)
            paragraph.paragraph_format.space_after = Pt(4)
            run = paragraph.add_run(token.content.rstrip("\n"))
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
        elif kind == "blockquote_open":
            quote_depth += 1
        elif kind == "blockquote_close":
            quote_depth = max(0, quote_depth - 1)
        elif kind in {"bullet_list_open", "ordered_list_open"}:
            list_stack.append({"ordered": kind == "ordered_list_open", "counter": 0})
        elif kind in {"bullet_list_close", "ordered_list_close"}:
            if list_stack:
                list_stack.pop()
        elif kind == "list_item_open":
            depth = len(list_stack)
            if list_stack and list_stack[-1]["ordered"]:
                list_stack[-1]["counter"] += 1
                pending_prefix = (f"{list_stack[-1]['counter']}. ", depth)
            else:
                pending_prefix = ("\u2022 ", depth)
        elif kind == "table_open":
            table_rows = []
        elif kind == "tr_open":
            current_row = []
        elif kind in {"th_open", "td_open"}:
            current_cell = []
        elif kind == "inline" and current_cell is not None:
            current_cell.append(token.children or [])
        elif kind in {"th_close", "td_close"}:
            if current_row is not None and current_cell is not None:
                current_row.append(current_cell)
            current_cell = None
        elif kind == "tr_close":
            if current_row is not None:
                table_rows.append(current_row)
            current_row = None
        elif kind == "table_close":
            if table_rows:
                columns = max(len(row) for row in table_rows)
                table = document.add_table(rows=len(table_rows), cols=columns)
                table.style = "Table Grid"
                for row_index, row in enumerate(table_rows):
                    for col_index in range(columns):
                        paragraph = table.cell(row_index, col_index).paragraphs[0]
                        cell_tokens = row[col_index][0] if col_index < len(row) else []
                        _add_inline(paragraph, cell_tokens, base_dir)
                        if row_index == 0:
                            for run in paragraph.runs:
                                run.bold = True
            table_rows = []
        elif kind == "hr":
            document.add_paragraph("———")
    document.save(str(destination))
def docx_to_txt(source: Path, destination: Path) -> None:
    document = Document(str(source))
    parts: list[str] = []
    try:
        blocks = document.iter_inner_content()
    except AttributeError:
        blocks = document.paragraphs
    for block in blocks:
        if hasattr(block, "rows"):
            for row in block.rows:
                parts.append("\t".join(cell.text.strip() for cell in row.cells))
            parts.append("")
        else:
            text = block.text.strip()
            if text:
                parts.append(text)
    write_text(destination, "\n\n".join(parts).strip() + "\n")


def _run_markdown_text(run) -> str:
    text = run.text
    if not text:
        return ""
    if run.bold:
        text = f"**{text}**"
    if run.italic:
        text = f"*{text}*"
    if run.font.name and run.font.name.lower() in {"consolas", "courier new"}:
        text = f"`{text}`"
    return text


def _table_to_markdown(table) -> str:
    rows: list[list[str]] = []
    for row in table.rows:
        values = [cell.text.replace("\n", "<br>").strip().replace("|", "\\|") for cell in row.cells]
        rows.append(values)
    if not rows:
        return ""
    columns = max(len(row) for row in rows)
    for row in rows:
        while len(row) < columns:
            row.append("")
    lines = ["| " + " | ".join(rows[0]) + " |", "| " + " | ".join(["---"] * columns) + " |"]
    for row in rows[1:]:
        lines.append("| " + " | ".join(row) + " |")
    return "\n".join(lines)


def _extract_images(document, assets_dir: Path) -> list[str]:
    saved: list[str] = []
    for part in document.part.related_parts.values():
        content_type = getattr(part, "content_type", "")
        if not content_type.startswith("image/"):
            continue
        extension = {
            "image/png": ".png",
            "image/jpeg": ".jpg",
            "image/gif": ".gif",
            "image/bmp": ".bmp",
            "image/tiff": ".tiff",
        }.get(content_type, ".bin")
        target = assets_dir / f"image{len(saved) + 1}{extension}"
        try:
            assets_dir.mkdir(parents=True, exist_ok=True)
            target.write_bytes(part.blob)
            saved.append(target.name)
        except Exception:
            continue
    return saved


def docx_to_md(source: Path, destination: Path) -> list[str]:
    document = Document(str(source))
    assets_dir = destination.parent / "assets"
    images = _extract_images(document, assets_dir)
    parts: list[str] = []
    try:
        blocks = document.iter_inner_content()
    except AttributeError:
        blocks = document.paragraphs
    for block in blocks:
        if hasattr(block, "rows"):
            markdown = _table_to_markdown(block)
            if markdown:
                parts.append(markdown)
            continue
        style_name = (block.style.name or "").lower()
        text = "".join(_run_markdown_text(run) for run in block.runs).strip()
        if not text:
            text = block.text.strip()
        if not text:
            continue
        if style_name.startswith("heading"):
            level = style_name.replace("heading", "").strip()
            parts.append(f"{'#' * (int(level) if level.isdigit() else 1)} {text}")
        elif style_name == "title":
            parts.append(f"# {text}")
        elif style_name.startswith("list bullet"):
            parts.append(f"- {text}")
        elif style_name.startswith("list number"):
            parts.append(f"1. {text}")
        elif style_name in {"quote", "intense quote"}:
            parts.append(f"> {text}")
        else:
            parts.append(text)
    if images:
        parts.append("")
        parts.append("## 文档图片")
        for name in images:
            parts.append(f"![{name}](assets/{name})")
    write_text(destination, "\n\n".join(parts).strip() + "\n")
    return images
