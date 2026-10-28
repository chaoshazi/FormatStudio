from __future__ import annotations

from copy import deepcopy
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from pptx import Presentation
from pptx.dml.color import RGBColor as PptRGB
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt as PptPt

from ..presets import PDF_BASE, PPTX_BASE, XLSX_BASE, DOCX_BASE
from .pdf_render import add_watermark


def default_style(applies_to: list[str], name: str = "自定义样式") -> dict[str, Any]:
    data = {
        "version": 1,
        "name": name,
        "applies_to": applies_to,
        "docx": deepcopy(DOCX_BASE),
        "xlsx": deepcopy(XLSX_BASE),
        "pptx": deepcopy(PPTX_BASE),
        "pdf": deepcopy(PDF_BASE),
    }
    return data


def normalize_style(data: dict[str, Any]) -> dict[str, Any]:
    result = default_style(data.get("applies_to") or ["docx"], data.get("name") or "自定义样式")
    result.update({key: value for key, value in data.items() if key in {"version", "name", "applies_to"}})
    for key in ("docx", "xlsx", "pptx", "pdf"):
        if isinstance(data.get(key), dict):
            result[key].update(data[key])
    pdf_value = data.get("pdf")
    if isinstance(pdf_value, dict) and isinstance(pdf_value.get("watermark"), dict):
        result["pdf"]["watermark"].update(pdf_value["watermark"])
    return result


def _font_name(docx_font: str, english: str) -> str:
    return docx_font or english or "微软雅黑"


def _set_docx_style(style, config: dict[str, Any], size_key: str, color_key: str = "text_color") -> None:
    font_cn = config.get("font_cn", "微软雅黑")
    font_en = config.get("font_en", "Calibri")
    style.font.name = _font_name(font_cn, font_en)
    if style.font.size is not None or size_key:
        style.font.size = Pt(float(config.get(size_key, 11)))
    color = str(config.get(color_key, "1F2937")).lstrip("#")
    style.font.color.rgb = RGBColor.from_string(color)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), font_cn)
    rfonts.set(qn("w:ascii"), font_en)
    rfonts.set(qn("w:hAnsi"), font_en)
    paragraph = style.paragraph_format
    paragraph.line_spacing = float(config.get("line_spacing", 1.5))
    paragraph.space_before = Pt(float(config.get("space_before", 0)))
    paragraph.space_after = Pt(float(config.get("space_after", 6)))


def _add_field(paragraph, instruction: str, placeholder: str = "1") -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(separate)
    run._r.append(text)
    run._r.append(end)


def apply_docx_style(path: Path, config: dict[str, Any]) -> list[str]:
    warnings: list[str] = []
    document = Document(str(path))
    background_color = str(config.get("page_background", "FFFFFF")).lstrip("#")
    existing = document.element.find(qn("w:background"))
    if existing is not None:
        document.element.remove(existing)
    background = OxmlElement("w:background")
    background.set(qn("w:color"), background_color)
    document.element.insert(0, background)
    font_cn = config.get("font_cn", "微软雅黑")
    font_en = config.get("font_en", "Calibri")
    style_map = {
        "Normal": "body_size",
        "Title": "title_size",
        "Heading 1": "h1_size",
        "Heading 2": "h2_size",
        "Heading 3": "h2_size",
    }
    for style_name, size_key in style_map.items():
        try:
            style = document.styles[style_name]
        except KeyError:
            continue
        color_key = "text_color" if style_name == "Normal" else "heading_color"
        _set_docx_style(style, config, size_key, color_key)
        if style_name == "Normal":
            style.paragraph_format.first_line_indent = Cm(float(config.get("first_line_indent", 0)))
    for section in document.sections:
        try:
            section.top_margin = Cm(float(config.get("margin_top_cm", 2.54)))
            section.bottom_margin = Cm(float(config.get("margin_bottom_cm", 2.54)))
            section.left_margin = Cm(float(config.get("margin_left_cm", 3.18)))
            section.right_margin = Cm(float(config.get("margin_right_cm", 3.18)))
            if config.get("orientation") == "landscape" and section.orientation != WD_ORIENT.LANDSCAPE:
                section.orientation = WD_ORIENT.LANDSCAPE
                if section.page_width < section.page_height:
                    section.page_width, section.page_height = section.page_height, section.page_width
            elif config.get("orientation", "portrait") == "portrait" and section.orientation != WD_ORIENT.PORTRAIT:
                section.orientation = WD_ORIENT.PORTRAIT
                if section.page_width > section.page_height:
                    section.page_width, section.page_height = section.page_height, section.page_width
        except Exception as exc:  # noqa: BLE001
            warnings.append(f"页面设置应用失败：{exc}")
        header_text = (config.get("header_text") or "").strip()
        footer_text = (config.get("footer_text") or "").strip()
        try:
            header = section.header
            header.is_linked_to_previous = False
            paragraph = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
            paragraph.text = header_text or "FORMAT STUDIO"
            paragraph.paragraph_format.space_after = Pt(6)
            header_fill = str(config.get("header_fill", config.get("primary_color", "2563EB"))).lstrip("#")
            p_pr = paragraph._p.get_or_add_pPr()
            shd = p_pr.find(qn("w:shd"))
            if shd is None:
                shd = OxmlElement("w:shd")
                p_pr.append(shd)
            shd.set(qn("w:fill"), header_fill)
            shd.set(qn("w:val"), "clear")
            for run in paragraph.runs:
                run.font.size = Pt(8.5)
                run.font.name = font_cn
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                rpr = run._element.get_or_add_rPr()
                rfonts = rpr.find(qn("w:rFonts"))
                if rfonts is None:
                    rfonts = rpr.makeelement(qn("w:rFonts"), {})
                    rpr.append(rfonts)
                rfonts.set(qn("w:eastAsia"), font_cn)
        except Exception as exc:  # noqa: BLE001
            warnings.append(f"页眉应用失败：{exc}")
        try:
            footer = section.footer
            footer.is_linked_to_previous = False
            paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
            paragraph.text = footer_text
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if config.get("page_number", True):
                if footer_text:
                    paragraph.add_run("    ")
                _add_field(paragraph, "PAGE")
            for run in paragraph.runs:
                run.font.size = Pt(9)
                run.font.name = font_cn
                rpr = run._element.get_or_add_rPr()
                rfonts = rpr.find(qn("w:rFonts"))
                if rfonts is None:
                    rfonts = rpr.makeelement(qn("w:rFonts"), {})
                    rpr.append(rfonts)
                rfonts.set(qn("w:eastAsia"), font_cn)
        except Exception as exc:  # noqa: BLE001
            warnings.append(f"页脚应用失败：{exc}")
    for table in document.tables:
        for row_index, row in enumerate(table.rows):
            for cell in row.cells:
                tc_pr = cell._tc.get_or_add_tcPr()
                shd = tc_pr.find(qn("w:shd"))
                if shd is None:
                    shd = OxmlElement("w:shd")
                    tc_pr.append(shd)
                if row_index == 0:
                    shd.set(qn("w:fill"), str(config.get("header_fill", "2563EB")).lstrip("#"))
                elif row_index % 2 == 0:
                    shd.set(qn("w:fill"), str(config.get("light_color", "EFF6FF")).lstrip("#"))
                shd.set(qn("w:val"), "clear")
                if row_index == 0:
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                            run.font.bold = True
    if config.get("force_direct_format"):
        body_size = Pt(float(config.get("body_size", 11)))
        for paragraph in document.paragraphs:
            style_name = (paragraph.style.name or "").lower()
            if style_name.startswith("heading") or style_name == "title":
                continue
            for run in paragraph.runs:
                run.font.name = _font_name(font_cn, font_en)
                run.font.size = body_size
                rpr = run._element.get_or_add_rPr()
                rfonts = rpr.find(qn("w:rFonts"))
                if rfonts is None:
                    rfonts = rpr.makeelement(qn("w:rFonts"), {})
                    rpr.append(rfonts)
                rfonts.set(qn("w:eastAsia"), font_cn)
    document.save(str(path))
    return warnings


def _border(color: str = "CBD5E1") -> Border:
    side = Side(style="thin", color=color)
    return Border(left=side, right=side, top=side, bottom=side)


def apply_xlsx_style(path: Path, config: dict[str, Any]) -> list[str]:
    warnings: list[str] = []
    workbook = load_workbook(path)
    try:
        font_name = config.get("font_name", "微软雅黑")
        font_size = float(config.get("font_size", 11))
        font_color = config.get("font_color", "0F172A")
        body_fill = config.get("body_fill", "")
        for worksheet in workbook.worksheets:
            max_row = max(1, worksheet.max_row)
            max_col = max(1, worksheet.max_column)
            for row_index in range(1, max_row + 1):
                for col_index in range(1, max_col + 1):
                    cell = worksheet.cell(row=row_index, column=col_index)
                    cell.font = Font(name=font_name, size=font_size, color=font_color)
                    cell.border = _border(config.get("border_color", "CBD5E1"))
                    if body_fill:
                        cell.fill = PatternFill("solid", fgColor=body_fill)
                    if row_index == 1:
                        cell.font = Font(
                            name=font_name,
                            size=font_size,
                            bold=bool(config.get("header_bold", True)),
                            color=str(config.get("header_color", "1E3A8A")),
                        )
                        cell.fill = PatternFill("solid", fgColor=str(config.get("header_fill", "DBEAFE")))
                        cell.alignment = Alignment(horizontal="center", vertical="center")
                    elif config.get("zebra_fill") and row_index % 2 == 0:
                        cell.fill = PatternFill("solid", fgColor=str(config.get("zebra_fill")))
            if config.get("freeze_header", True):
                worksheet.sheet_properties.tabColor = str(config.get("tab_color", config.get("header_fill", "2563EB"))).lstrip("#")
                worksheet.freeze_panes = "A2"
            if config.get("auto_width", True):
                for col_index in range(1, max_col + 1):
                    letter = get_column_letter(col_index)
                    width = 8
                    for row_index in range(1, min(max_row, 80) + 1):
                        value = worksheet.cell(row=row_index, column=col_index).value
                        text = "" if value is None else str(value)
                        width = max(width, min(50, sum(2 if ord(ch) > 127 else 1 for ch in text) + 2))
                    worksheet.column_dimensions[letter].width = width
            if config.get("landscape", True):
                worksheet.page_setup.orientation = "landscape"
            if config.get("fit_to_width", True):
                worksheet.page_setup.fitToWidth = 1
                worksheet.sheet_properties.pageSetUpPr.fitToPage = True
            worksheet.print_title_rows = "1:1"
    except Exception as exc:  # noqa: BLE001
        warnings.append(f"Excel 样式应用部分失败：{exc}")
    finally:
        workbook.save(str(path))
        workbook.close()
    return warnings


def _iter_pptx_shapes(shapes):
    for shape in shapes:
        yield shape
        if hasattr(shape, "shapes"):
            yield from _iter_pptx_shapes(shape.shapes)


def _apply_pptx_theme(presentation, config: dict[str, Any]) -> None:
    from lxml import etree
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT

    theme_part = presentation.slide_masters[0].part.part_related_by(RT.THEME)
    root = etree.fromstring(theme_part.blob)
    namespace = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    fonts = [str(config.get("title_font", "????")), str(config.get("body_font", "????"))]
    for scheme, font in zip(("majorFont", "minorFont"), fonts):
        for tag in ("latin", "ea", "cs"):
            nodes = root.findall(f".//a:{scheme}/a:{tag}", namespaces=namespace)
            for node in nodes:
                node.set("typeface", font)
    colors = config.get("theme_colors") or []
    if len(colors) >= 6:
        mapping = {
            "dk1": colors[1],
            "lt1": colors[4],
            "dk2": colors[0],
            "lt2": colors[3],
            "accent1": colors[0],
            "accent2": colors[2],
            "accent3": colors[5],
            "accent4": colors[1],
            "accent5": colors[3],
            "accent6": colors[4],
        }
        for key, value in mapping.items():
            nodes = root.findall(f".//a:clrScheme/a:{key}/a:srgbClr", namespaces=namespace)
            for node in nodes:
                node.set("val", str(value).lstrip("#"))
    theme_part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def apply_pptx_style(path: Path, config: dict[str, Any]) -> list[str]:
    warnings: list[str] = []
    presentation = Presentation(str(path))
    try:
        _apply_pptx_theme(presentation, config)
        if config.get("slide_width_inches") and config.get("slide_height_inches"):
            presentation.slide_width = PptPt(float(config["slide_width_inches"]) * 72)
            presentation.slide_height = PptPt(float(config["slide_height_inches"]) * 72)
        background = str(config.get("background_color", "F8FAFC")).lstrip("#")
        primary = str(config.get("primary_color", "2563EB")).lstrip("#")
        text_color = str(config.get("text_color", "0F172A")).lstrip("#")
        for slide in presentation.slides:
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = PptRGB.from_string(background)
            for shape in list(slide.shapes):
                if shape.name == "FormatStudioAccent":
                    shape._element.getparent().remove(shape._element)
            accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, presentation.slide_width, Inches(0.07))
            accent.name = "FormatStudioAccent"
            accent.fill.solid()
            accent.fill.fore_color.rgb = PptRGB.from_string(primary)
            accent.line.fill.background()
            title_shape = slide.shapes.title
            for shape in _iter_pptx_shapes(slide.shapes):
                if not getattr(shape, "has_text_frame", False):
                    continue
                font_name = config.get("title_font" if shape is title_shape else "body_font", "????")
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        run.font.name = font_name
                        try:
                            if run.font.color.type is not None and str(run.font.color.rgb).upper() in {"FFFFFF", "F8FAFC", "F9FAFB"}:
                                run.font.color.rgb = PptRGB.from_string(text_color)
                        except Exception:
                            pass
                if shape is title_shape:
                    for paragraph in shape.text_frame.paragraphs:
                        for run in paragraph.runs:
                            if run.font.size is None:
                                run.font.size = PptPt(30)
    except Exception as exc:  # noqa: BLE001
        warnings.append(f"PPT 样式应用部分失败：{exc}")
    presentation.save(str(path))
    return warnings


def apply_pdf_watermark(path: Path, config: dict[str, Any], password: str | None = None) -> list[str]:
    watermark = config.get("watermark") or {}
    if not watermark.get("enabled") or not str(watermark.get("text", "")).strip():
        return []
    temp = path.with_name(path.stem + ".watermark" + path.suffix)
    add_watermark(
        path,
        temp,
        text=str(watermark.get("text", "")),
        opacity=float(watermark.get("opacity", 0.15)),
        angle=int(watermark.get("angle", 45)),
        size=int(watermark.get("size", 48)),
        color=str(watermark.get("color", "94A3B8")),
        tile=bool(watermark.get("tile", False)),
        password=password,
    )
    temp.replace(path)
    return []


def apply_style(path: Path, fmt: str, data: dict[str, Any], password: str | None = None) -> list[str]:
    data = normalize_style(data)
    if fmt == "docx":
        return apply_docx_style(path, data["docx"])
    if fmt == "xlsx":
        return apply_xlsx_style(path, data["xlsx"])
    if fmt == "pptx":
        return apply_pptx_style(path, data["pptx"])
    if fmt == "pdf":
        return apply_pdf_watermark(path, data["pdf"], password)
    return []


def import_style_from_sample(path: Path, fmt: str, name: str) -> tuple[dict[str, Any], list[str]]:
    warnings: list[str] = []
    data = default_style([fmt], name)
    if fmt == "docx":
        document = Document(str(path))
        config = data["docx"]
        for style_name, size_key in {"Normal": "body_size", "Title": "title_size", "Heading 1": "h1_size", "Heading 2": "h2_size"}.items():
            try:
                style = document.styles[style_name]
                if style.font.name:
                    config["font_en"] = style.font.name
                if style.font.size:
                    config[size_key] = round(style.font.size.pt, 1)
                if style.paragraph_format.line_spacing:
                    config["line_spacing"] = float(style.paragraph_format.line_spacing)
            except Exception:
                warnings.append(f"未能读取 {style_name} 样式")
        if document.sections:
            section = document.sections[0]
            config["margin_top_cm"] = round(section.top_margin.cm, 2)
            config["margin_bottom_cm"] = round(section.bottom_margin.cm, 2)
            config["margin_left_cm"] = round(section.left_margin.cm, 2)
            config["margin_right_cm"] = round(section.right_margin.cm, 2)
            config["orientation"] = "landscape" if section.page_width > section.page_height else "portrait"
            try:
                config["header_text"] = section.header.paragraphs[0].text
            except Exception:
                pass
            try:
                config["footer_text"] = section.footer.paragraphs[0].text
            except Exception:
                pass
        try:
            normal = document.styles["Normal"]
            rpr = normal.element.find(qn("w:rPr"))
            if rpr is not None:
                rfonts = rpr.find(qn("w:rFonts"))
                if rfonts is not None and rfonts.get(qn("w:eastAsia")):
                    config["font_cn"] = rfonts.get(qn("w:eastAsia"))
        except Exception:
            pass
    elif fmt == "xlsx":
        workbook = load_workbook(path)
        try:
            worksheet = workbook.worksheets[0]
            config = data["xlsx"]
            header = worksheet.cell(row=1, column=1)
            if header.font:
                config["font_name"] = header.font.name or config["font_name"]
                config["font_size"] = header.font.size or config["font_size"]
                config["header_bold"] = bool(header.font.bold)
            try:
                if header.fill and header.fill.fgColor and header.fill.fgColor.rgb:
                    rgb = str(header.fill.fgColor.rgb)[-6:]
                    config["header_fill"] = rgb
                if header.font and header.font.color and header.font.color.rgb:
                    config["header_color"] = str(header.font.color.rgb)[-6:]
            except Exception:
                pass
            if worksheet.freeze_panes:
                config["freeze_header"] = worksheet.freeze_panes == "A2"
            config["landscape"] = worksheet.page_setup.orientation == "landscape"
        finally:
            workbook.close()
    elif fmt == "pptx":
        presentation = Presentation(str(path))
        config = data["pptx"]
        config["slide_width_inches"] = round(presentation.slide_width / 914400, 2)
        config["slide_height_inches"] = round(presentation.slide_height / 914400, 2)
        try:
            from pptx.opc.constants import RELATIONSHIP_TYPE as RT

            from lxml import etree

            theme_part = presentation.slide_masters[0].part.part_related_by(RT.THEME)
            theme_root = etree.fromstring(theme_part.blob)
            namespace = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
            fonts = theme_root.findall(".//a:latin", namespaces=namespace)
            if fonts:
                config["title_font"] = fonts[0].get("typeface") or config["title_font"]
            if len(fonts) > 1:
                config["body_font"] = fonts[1].get("typeface") or config["body_font"]
        except Exception as exc:  # noqa: BLE001
            warnings.append(f"未能读取 PPT 主题字体：{exc}")
    else:
        raise ValueError("仅支持从 docx/xlsx/pptx 导入样式")
    return normalize_style(data), warnings
