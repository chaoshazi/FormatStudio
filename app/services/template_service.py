from __future__ import annotations

import json
import re
import shutil
from pathlib import Path
from typing import Any

from docx import Document
from docx.shared import Pt
from openpyxl import load_workbook
from pptx import Presentation
from pptx.util import Pt as PptPt

from ..formats import safe_stem
from ..utils import read_csv_rows, unique_path, write_text, zip_paths
from .styles_service import apply_style

PLACEHOLDER = re.compile(r"\{\{\s*([A-Za-z0-9_\u4e00-\u9fff]+)\s*\}\}")
JINJA_FORBIDDEN = re.compile(r"\{%|\{#|\{\{\s*[^}]*[\+\-\*/\(\)\[\]\.]")


class TemplateError(RuntimeError):
    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code
        self.message = message


def _count_fields(texts: list[str], found: dict[str, dict[str, Any]]) -> None:
    for text in texts:
        for match in PLACEHOLDER.finditer(text or ""):
            field = match.group(1)
            entry = found.setdefault(field, {"name": field, "count": 0, "sample": ""})
            entry["count"] += 1
            if not entry["sample"]:
                start = max(0, match.start() - 20)
                end = min(len(text), match.end() + 20)
                entry["sample"] = text[start:end].strip()


def _docx_texts(document) -> list[str]:
    texts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                texts.extend(paragraph.text for paragraph in cell.paragraphs)
    for section in document.sections:
        for container in (section.header, section.footer):
            texts.extend(paragraph.text for paragraph in container.paragraphs)
    return texts


def _xlsx_texts(workbook) -> list[str]:
    texts = []
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    texts.append(cell.value)
        for name in ("oddHeader", "oddFooter"):
            container = getattr(worksheet, name, None)
            if container and getattr(container, "center", None):
                texts.append(container.center)
    return texts


def _pptx_texts(presentation) -> list[str]:
    texts: list[str] = []

    def walk(shapes):
        for shape in shapes:
            if getattr(shape, "has_text_frame", False):
                texts.extend(paragraph.text for paragraph in shape.text_frame.paragraphs)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    for cell in row.cells:
                        texts.extend(paragraph.text for paragraph in cell.text_frame.paragraphs)
            if hasattr(shape, "shapes"):
                walk(shape.shapes)

    for slide in presentation.slides:
        walk(slide.shapes)
    return texts


def scan_template(path: Path, fmt: str) -> list[dict[str, Any]]:
    found: dict[str, dict[str, Any]] = {}
    if fmt == "docx":
        document = Document(str(path))
        texts = _docx_texts(document)
    elif fmt == "xlsx":
        workbook = load_workbook(path, read_only=True)
        try:
            texts = _xlsx_texts(workbook)
        finally:
            workbook.close()
    elif fmt == "pptx":
        presentation = Presentation(str(path))
        texts = _pptx_texts(presentation)
    else:
        raise TemplateError("INVALID_TEMPLATE", "模板仅支持 docx、xlsx、pptx")
    for text in texts:
        if JINJA_FORBIDDEN.search(text or ""):
            raise TemplateError(
                "INVALID_TEMPLATE",
                f"模板含不支持的 Jinja 语法，仅允许 {{{{字段名}}}}：{text[:80]}",
            )
    _count_fields(texts, found)
    return sorted(found.values(), key=lambda item: item["name"])


def _replace_in_paragraph(paragraph, values: dict[str, str], missing: set[str]) -> None:
    runs = paragraph.runs
    if not runs:
        return
    full = "".join(run.text for run in runs)
    matches = list(PLACEHOLDER.finditer(full))
    if not matches:
        return
    starts: list[int] = []
    ends: list[int] = []
    offset = 0
    for run in runs:
        starts.append(offset)
        offset += len(run.text)
        ends.append(offset)
    for match in reversed(matches):
        field = match.group(1)
        if field not in values:
            missing.add(field)
            continue
        replacement = str(values[field])
        start_run = next((index for index, end in enumerate(ends) if end > match.start()), None)
        end_run = next((index for index, start in enumerate(starts) if start >= match.end()), len(runs) - 1)
        if start_run is None:
            continue
        if end_run < start_run:
            end_run = start_run
        start_offset = match.start() - starts[start_run]
        end_offset = match.end() - starts[end_run]
        if start_run == end_run:
            runs[start_run].text = runs[start_run].text[:start_offset] + replacement + runs[start_run].text[end_offset:]
        else:
            runs[start_run].text = runs[start_run].text[:start_offset] + replacement
            for index in range(start_run + 1, end_run):
                runs[index].text = ""
            runs[end_run].text = runs[end_run].text[end_offset:]


def fill_docx(path: Path, values: dict[str, str]) -> set[str]:
    document = Document(str(path))
    missing: set[str] = set()
    for paragraph in document.paragraphs:
        _replace_in_paragraph(paragraph, values, missing)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    _replace_in_paragraph(paragraph, values, missing)
    for section in document.sections:
        for container in (section.header, section.footer):
            for paragraph in container.paragraphs:
                _replace_in_paragraph(paragraph, values, missing)
    document.save(str(path))
    return missing


def fill_xlsx(path: Path, values: dict[str, str]) -> set[str]:
    workbook = load_workbook(path)
    missing: set[str] = set()
    try:
        for worksheet in workbook.worksheets:
            for row in worksheet.iter_rows():
                for cell in row:
                    if not isinstance(cell.value, str):
                        continue
                    text = cell.value
                    new_text = text
                    for match in reversed(list(PLACEHOLDER.finditer(text))):
                        field = match.group(1)
                        if field not in values:
                            missing.add(field)
                            continue
                        new_text = new_text[: match.start()] + str(values[field]) + new_text[match.end():]
                    cell.value = new_text
        workbook.save(str(path))
    finally:
        workbook.close()
    return missing


def _fill_pptx_shapes(shapes, values: dict[str, str], missing: set[str]) -> None:
    for shape in shapes:
        if getattr(shape, "has_text_frame", False):
            for paragraph in shape.text_frame.paragraphs:
                _replace_in_paragraph(paragraph, values, missing)
        if getattr(shape, "has_table", False):
            for row in shape.table.rows:
                for cell in row.cells:
                    for paragraph in cell.text_frame.paragraphs:
                        _replace_in_paragraph(paragraph, values, missing)
        if hasattr(shape, "shapes"):
            _fill_pptx_shapes(shape.shapes, values, missing)


def fill_pptx(path: Path, values: dict[str, str]) -> set[str]:
    presentation = Presentation(str(path))
    missing: set[str] = set()
    for slide in presentation.slides:
        _fill_pptx_shapes(slide.shapes, values, missing)
    presentation.save(str(path))
    return missing


def fill_template(path: Path, fmt: str, values: dict[str, str]) -> set[str]:
    if fmt == "docx":
        return fill_docx(path, values)
    if fmt == "xlsx":
        return fill_xlsx(path, values)
    if fmt == "pptx":
        return fill_pptx(path, values)
    raise TemplateError("INVALID_TEMPLATE", f"不支持的模板格式：{fmt}")


def parse_data(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        headers, rows, _ = read_csv_rows(path)
        return headers, rows
    if suffix == ".json":
        try:
            payload = json.loads(path.read_text(encoding="utf-8-sig"))
        except Exception as exc:
            raise TemplateError("INVALID_DATA", f"JSON 解析失败：{exc}") from exc
        if isinstance(payload, dict) and isinstance(payload.get("rows"), list):
            payload = payload["rows"]
        if not isinstance(payload, list):
            raise TemplateError("INVALID_DATA", "JSON 数据应为对象数组，或包含 rows 数组")
        headers: list[str] = []
        rows: list[dict[str, str]] = []
        for item in payload:
            if not isinstance(item, dict):
                raise TemplateError("INVALID_DATA", "JSON 数组中的每一项都必须是对象")
            record: dict[str, str] = {}
            for key, value in item.items():
                if isinstance(value, (dict, list)):
                    raise TemplateError("INVALID_DATA", f"字段 {key} 是嵌套对象，暂不支持")
                if key not in headers:
                    headers.append(key)
                record[key] = "" if value is None else str(value)
            rows.append(record)
        return headers, rows
    if suffix == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            worksheet = workbook.worksheets[0]
            values = list(worksheet.iter_rows(values_only=True))
            if not values:
                return [], []
            headers = [str(value).strip() if value is not None else f"字段{index + 1}" for index, value in enumerate(values[0])]
            if len(set(headers)) != len(headers):
                raise TemplateError("INVALID_DATA", "数据表头存在重复列名")
            rows = []
            for raw in values[1:]:
                if not any(value is not None and str(value).strip() for value in raw):
                    continue
                record = {}
                for index, header in enumerate(headers):
                    value = raw[index] if index < len(raw) else None
                    record[header] = "" if value is None else str(value)
                rows.append(record)
            return headers, rows
        finally:
            workbook.close()
    raise TemplateError("INVALID_DATA", "数据源仅支持 CSV、JSON、XLSX")


def _output_name(template_stem: str, row: dict[str, str], index: int, suffix: str) -> str:
    requested = row.get("__filename__") or row.get("文件名")
    if requested:
        stem = safe_stem(requested)
    else:
        stem = f"{safe_stem(template_stem)}_{index:03d}"
    return f"{stem}{suffix}"


def generate_from_template(
    template_path: Path,
    template_format: str,
    data_path: Path,
    output_dir: Path,
    style_data: dict[str, Any] | None = None,
) -> dict[str, Any]:
    headers, rows = parse_data(data_path)
    if not rows:
        raise TemplateError("INVALID_DATA", "数据源没有可用于生成的行")
    scanned = [field["name"] for field in scan_template(template_path, template_format)]
    missing_headers = [field for field in scanned if field not in headers and field not in {"__filename__", "文件名"}]
    if missing_headers:
        raise TemplateError("MISSING_FIELD", f"数据源缺少字段：{', '.join(missing_headers)}")
    output_dir.mkdir(parents=True, exist_ok=True)
    outputs: list[Path] = []
    warnings: list[str] = []
    used_names: set[str] = set()
    for index, row in enumerate(rows, start=1):
        target = unique_path(output_dir / _output_name(template_path.stem, row, index, template_path.suffix.lower()))
        if target.name in used_names:
            target = unique_path(target)
        used_names.add(target.name)
        shutil.copy2(template_path, target)
        missing = fill_template(target, template_format, row)
        if missing:
            target.unlink(missing_ok=True)
            warnings.append(f"第 {index} 行缺少字段：{', '.join(sorted(missing))}")
            continue
        if style_data:
            try:
                apply_style(target, template_format, style_data)
            except Exception as exc:  # noqa: BLE001
                warnings.append(f"第 {index} 行样式应用失败：{exc}")
        outputs.append(target)
    if not outputs:
        raise TemplateError("MISSING_FIELD", "所有数据行都缺少必填字段，未生成文件")
    archive = None
    if len(outputs) > 1:
        archive = output_dir / f"{safe_stem(template_path.stem)}_生成结果.zip"
        zip_paths(outputs, archive)
    return {
        "outputs": outputs,
        "archive": archive,
        "rows": len(rows),
        "generated": len(outputs),
        "warnings": warnings,
        "headers": headers,
    }


def create_sample_document(fmt: str, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    if fmt == "docx":
        document = Document()
        document.add_heading("样式预览标题", level=1)
        document.add_paragraph("这是一段用于预览样式的正文。字体、字号、行距、段距与页边距都会随样式模板变化。")
        document.add_heading("二级标题", level=2)
        document.add_paragraph("第二段正文用于观察段落间距与首行缩进。")
        table = document.add_table(rows=2, cols=3)
        table.style = "Table Grid"
        for index, value in enumerate(["项目", "数量", "备注"]):
            table.cell(0, index).text = value
        for index, value in enumerate(["A", "12", "示例"]):
            table.cell(1, index).text = value
        document.save(str(path))
    elif fmt == "xlsx":
        import openpyxl

        workbook = openpyxl.Workbook()
        worksheet = workbook.active
        worksheet.title = "示例数据"
        headers = ["项目", "数量", "状态", "备注"]
        worksheet.append(headers)
        for row in (["需求梳理", 12, "完成", "示例"], ["接口联调", 8, "进行中", "示例"], ["验收", 4, "待开始", ""]):
            worksheet.append(row)
        workbook.save(str(path))
    elif fmt == "pptx":
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[0])
        slide.shapes.title.text = "样式预览"
        slide.placeholders[1].text = "标题字体、正文字体与幻灯片尺寸会随样式模板变化。"
        slide2 = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide2.shapes.title.text = "内容页"
        slide2.placeholders[1].text = "示例正文\n第二行"
        presentation.save(str(path))
    else:
        write_text(path, "示例文本\n")
    return path