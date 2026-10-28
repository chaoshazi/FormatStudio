from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.formatting.rule import CellIsRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

THIN = Side(style="thin", color="E2E8F0")


def fill(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=color.lstrip("#"))


def font(size: float = 10, bold: bool = False, color: str = "0F172A", name: str = "微软雅黑") -> Font:
    return Font(name=name, size=size, bold=bold, color=color.lstrip("#"))


def border(color: str = "E2E8F0") -> Border:
    side = Side(style="thin", color=color.lstrip("#"))
    return Border(left=side, right=side, top=side, bottom=side)


def text(
    ws,
    row: int,
    col: int,
    value: Any,
    *,
    size: float = 10,
    bold: bool = False,
    color: str = "0F172A",
    bg: str | None = None,
    align: str = "left",
    merge_to: int | None = None,
    wrap: bool = False,
    border_color: str | None = None,
    height: float | None = None,
) -> None:
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = font(size, bold, color)
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap, indent=1 if align == "left" else 0)
    if bg:
        cell.fill = fill(bg)
    if border_color:
        cell.border = border(border_color)
    if merge_to and merge_to > col:
        ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=merge_to)
    if height:
        ws.row_dimensions[row].height = height


def build_data_sheet(workbook: Workbook, spec: dict[str, Any], palette: dict[str, str]):
    ws = workbook.create_sheet(spec["sheet_name"])
    ws.sheet_view.showGridLines = False
    headers = spec["headers"]
    last_col = len(headers)
    for col, width in enumerate((16, 16, 16, 16, 14, 14, 14, 14), start=1):
        if col <= max(6, last_col):
            ws.column_dimensions[get_column_letter(col)].width = width
    text(ws, 1, 1, spec["title"], size=16, bold=True, color=palette["white"], bg=palette["dark"], merge_to=last_col, height=34)
    text(ws, 2, 1, spec["subtitle"], size=9.5, color=palette["muted"], bg=palette["light"], merge_to=last_col, height=22)
    for index, (label, value) in enumerate(spec["kpis"][:3]):
        start = 1 + index * 2 if last_col >= 6 else 1 + index
        label_col = start
        value_col = start + 1 if last_col >= 6 else start
        text(ws, 3, label_col, label, size=8.5, bold=True, color=palette["muted"], bg=palette["light"], align="center", border_color=palette["border"])
        text(ws, 3, value_col, value, size=10, bold=True, color=palette["primary"], bg=palette["white"], align="center", border_color=palette["border"])
    for col_index, header in enumerate(headers, start=1):
        text(ws, 4, col_index, header, size=9.5, bold=True, color=palette["white"], bg=palette["primary"], align="center", border_color=palette["border"], height=24)
    for row_index, row in enumerate(spec["rows"], start=5):
        for col_index, value in enumerate(row, start=1):
            cell = ws.cell(row=row_index, column=col_index, value=value)
            cell.font = font(9.5, False, palette["dark"])
            cell.border = border(palette["border"])
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.fill = fill(palette["white"] if row_index % 2 else palette["light"])
            header = headers[col_index - 1]
            if "率" in header or "进度" in header or "达成" in header:
                cell.number_format = "0%"
            elif any(marker in header for marker in ("金额", "预算", "支出", "单价", "销售", "预计", "差异")):
                cell.number_format = "#,##0"
    last_row = 4 + len(spec["rows"])
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:{get_column_letter(last_col)}{last_row}"
    ws.print_title_rows = "4:4"
    ws.print_area = f"A1:{get_column_letter(last_col)}{last_row}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.sheet_properties.tabColor = palette["primary"]
    ws.oddFooter.center.text = "Format Studio"
    ws.oddFooter.center.size = 8
    ws.oddFooter.center.color = palette["muted"]
    for col_index, header in enumerate(headers, start=1):
        letter = get_column_letter(col_index)
        if any(marker in header for marker in ("率", "进度", "达成")):
            ws.conditional_formatting.add(
                f"{letter}5:{letter}{last_row}",
                DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color=palette["accent"]),
            )
        if "状态" in header:
            ws.conditional_formatting.add(
                f"{letter}5:{letter}{last_row}",
                CellIsRule(operator="equal", formula=['"已完成"'], fill=fill("DCFCE7")),
            )
            ws.conditional_formatting.add(
                f"{letter}5:{letter}{last_row}",
                CellIsRule(operator="equal", formula=['"进行中"'], fill=fill("DBEAFE")),
            )
            ws.conditional_formatting.add(
                f"{letter}5:{letter}{last_row}",
                CellIsRule(operator="equal", formula=['"未开始"'], fill=fill("F1F5F9")),
            )
    chart = None
    columns = spec.get("chart_columns") or []
    if spec.get("chart") and columns:
        chart = BarChart() if spec.get("chart") != "pie" else PieChart()
        chart.title = spec["sheet_name"]
        chart.height = 8
        chart.width = 15
        chart.style = 10
        chart.add_data(Reference(ws, min_col=min(columns), max_col=max(columns), min_row=4, max_row=last_row), titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=1, min_row=5, max_row=last_row))
        ws.add_chart(chart, f"{get_column_letter(last_col + 2)}4")
    return ws, chart
def add_chart(ws, data_ws, spec, palette: dict[str, str], anchor: str) -> None:
    chart = BarChart() if spec.get("chart") != "pie" else PieChart()
    chart.title = spec["sheet_name"]
    chart.height = 8.5
    chart.width = 16
    chart.style = 10
    last_col = len(spec["headers"])
    last_row = 4 + len(spec["rows"])
    columns = spec.get("chart_columns") or []
    if not columns:
        return
    chart.add_data(Reference(data_ws, min_col=min(columns), max_col=max(columns), min_row=4, max_row=last_row), titles_from_data=True)
    chart.set_categories(Reference(data_ws, min_col=1, min_row=5, max_row=last_row))
    if chart.series:
        chart.series[0].graphicalProperties.solidFill = palette["primary"]
        chart.series[0].graphicalProperties.line.solidFill = palette["primary"]
        if len(chart.series) > 1:
            chart.series[1].graphicalProperties.solidFill = palette["accent"]
    ws.add_chart(chart, anchor)


def build_cover(workbook: Workbook, spec: dict[str, Any], palette: dict[str, str], variant: str, data_ws) -> None:
    cover = workbook.create_sheet("概览", 0)
    cover.sheet_view.showGridLines = False
    cover.sheet_properties.tabColor = palette["dark"]
    for col, width in zip("ABCDEF", (15, 15, 15, 15, 15, 15)):
        cover.column_dimensions[col].width = width
    cover.oddFooter.center.text = "Format Studio"
    cover.oddFooter.center.size = 8
    cover.oddFooter.center.color = palette["muted"]
    if variant == "executive":
        cover.merge_cells("A1:B29")
        cell = cover["A1"]
        cell.fill = fill(palette["dark"])
        cell.value = "FORMAT STUDIO\n\n" + spec["name"].split(" · ")[0] + "\n\n{{company}}\n{{date}}\n\n" + "0" + str(1)
        cell.font = font(13, True, palette["white"])
        cell.alignment = Alignment(vertical="center", horizontal="left", wrap_text=True, indent=2)
        text(cover, 1, 3, "EXECUTIVE BRIEF / {{date}}", size=9, bold=True, color=palette["primary"], merge_to=6, height=24)
        text(cover, 2, 3, spec["title"], size=26, bold=True, color=palette["dark"], merge_to=6, height=46)
        text(cover, 3, 3, spec["subtitle"], size=11, color=palette["muted"], merge_to=6, height=26)
        for index, (label, value) in enumerate(spec["meta"][:4]):
            row = 5 + index // 2
            col = 3 + (index % 2) * 2
            text(cover, row, col, label, size=8.5, bold=True, color=palette["muted"], bg=palette["light"], border_color=palette["border"])
            text(cover, row, col + 1, value, size=9.5, color=palette["dark"], border_color=palette["border"])
        for index, (label, value) in enumerate(spec["kpis"][:2]):
            col = 3 + index * 2
            cover.merge_cells(start_row=9, start_column=col, end_row=9, end_column=col + 1)
            text(cover, 9, col, label, size=8.5, color=palette["muted"], bg=palette["light"], align="center", border_color=palette["border"])
            cover.merge_cells(start_row=10, start_column=col, end_row=12, end_column=col + 1)
            text(cover, 10, col, value, size=18, bold=True, color=palette["primary"], align="center", border_color=palette["border"])
        cover.merge_cells("C15:F20")
        note = cover["C15"]
        note.value = spec.get("notes", "")
        note.font = font(10, False, palette["dark"])
        note.fill = fill(palette["light"])
        note.alignment = Alignment(wrap_text=True, vertical="center", indent=2)
        cover.print_area = "A1:F29"
    elif variant == "tracker":
        text(cover, 1, 1, spec["title"], size=22, bold=True, color=palette["white"], bg=palette["primary"], merge_to=6, height=42)
        text(cover, 2, 1, spec["subtitle"], size=10.5, color=palette["dark"], bg=palette["light"], merge_to=6, height=28)
        for index, (label, value) in enumerate(spec["kpis"][:3]):
            col = 1 + index * 2
            cover.merge_cells(start_row=4, start_column=col, end_row=4, end_column=col + 1)
            text(cover, 4, col, label, size=8.5, bold=True, color=palette["muted"], bg=palette["light"], align="center", border_color=palette["border"])
            cover.merge_cells(start_row=5, start_column=col, end_row=6, end_column=col + 1)
            text(cover, 5, col, value, size=18, bold=True, color=palette["primary"], align="center", border_color=palette["border"])
        for index, (label, value) in enumerate(spec["meta"][:4]):
            row = 8 + index // 2
            col = 1 + (index % 2) * 3
            text(cover, row, col, label, size=8.5, bold=True, color=palette["muted"], bg=palette["light"], border_color=palette["border"])
            text(cover, row, col + 1, value, size=9.5, color=palette["dark"], border_color=palette["border"], merge_to=col + 2)
        text(cover, 11, 1, "状态图例", size=9, bold=True, color=palette["dark"], merge_to=6)
        for index, (label, color) in enumerate((("已完成", "DCFCE7"), ("进行中", "DBEAFE"), ("未开始", "F1F5F9"))):
            text(cover, 12, 1 + index * 2, label, size=9, color=palette["dark"], bg=color, align="center", merge_to=2 + index * 2, border_color=palette["border"])
        text(cover, 14, 1, spec["sheet_name"], size=11, bold=True, color=palette["primary"], merge_to=6)
        for row_index, row in enumerate(spec["rows"][:4], start=15):
            for col_index, value in enumerate(row[:6], start=1):
                if isinstance(value, str) and value.startswith("="):
                    value = "—"
                bg = palette["primary"] if row_index == 15 else (palette["white"] if row_index % 2 else palette["light"])
                color = palette["white"] if row_index == 15 else palette["dark"]
                text(cover, row_index, col_index, value, size=8.5, bold=row_index == 15, color=color, bg=bg, border_color=palette["border"], wrap=True)
        cover.merge_cells("A20:F22")
        note = cover["A20"]
        note.value = spec.get("notes", "")
        note.font = font(9, False, palette["muted"])
        note.alignment = Alignment(wrap_text=True, vertical="top", indent=1)
        cover.print_area = "A1:F22"
    else:
        text(cover, 1, 1, "DASHBOARD / {{date}}", size=9, bold=True, color=palette["accent"], bg=palette["dark"], merge_to=6, height=24)
        text(cover, 2, 1, spec["title"], size=24, bold=True, color=palette["white"], bg=palette["dark"], merge_to=6, height=42)
        text(cover, 3, 1, spec["subtitle"], size=10.5, color=palette["muted"], bg=palette["light"], merge_to=6, height=26)
        for index, (label, value) in enumerate(spec["kpis"][:3]):
            col = 1 + index * 2
            cover.merge_cells(start_row=5, start_column=col, end_row=5, end_column=col + 1)
            text(cover, 5, col, label, size=8.5, bold=True, color=palette["muted"], bg=palette["light"], align="center", border_color=palette["border"])
            cover.merge_cells(start_row=6, start_column=col, end_row=7, end_column=col + 1)
            text(cover, 6, col, value, size=20, bold=True, color=palette["primary"], align="center", border_color=palette["border"])
        for index, (label, value) in enumerate(spec["meta"][:4]):
            row = 9 + index // 2
            col = 1 + (index % 2) * 3
            text(cover, row, col, label, size=8.5, bold=True, color=palette["muted"], bg=palette["light"], border_color=palette["border"])
            text(cover, row, col + 1, value, size=9.5, color=palette["dark"], border_color=palette["border"], merge_to=col + 2)
        add_chart(cover, data_ws, spec, palette, "A12")
        cover.print_area = "A1:F31"
    cover.page_setup.orientation = "portrait"
    cover.page_setup.fitToWidth = 1
    cover.page_setup.fitToHeight = 1
    cover.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)


def build_guide(workbook: Workbook, spec: dict[str, Any], palette: dict[str, str]) -> None:
    guide = workbook.create_sheet("使用说明")
    guide.sheet_view.showGridLines = False
    guide.column_dimensions["A"].width = 100
    text(guide, 1, 1, "模板使用说明", size=16, bold=True, color=palette["primary"], height=32)
    rules = [
        "1. 在数据源中为模板里的双花括号占位符提供对应值，一行数据生成一份文件。",
        "2. 示例数据用于展示版式，可在生成后替换为实际业务数据。",
        f"3. {spec.get('notes', '')}",
    ]
    for row, rule in enumerate(rules, start=3):
        text(guide, row, 1, rule, size=10, color=palette["dark"], wrap=True, height=28)
    guide.print_area = f"A1:A{2 + len(rules)}"
    guide.page_setup.orientation = "portrait"
    guide.page_setup.fitToWidth = 1
    guide.page_setup.fitToHeight = 1
    guide.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    guide.sheet_properties.tabColor = palette["accent"]


def build_xlsx_rich(spec: dict[str, Any], path: Path, palettes: dict[str, dict[str, str]], variant: str) -> None:
    palette = palettes[spec["palette"]]
    workbook = Workbook()
    default = workbook.active
    workbook.remove(default)
    data_ws, _ = build_data_sheet(workbook, spec, palette)
    build_cover(workbook, spec, palette, variant, data_ws)
    build_guide(workbook, spec, palette)
    workbook.active = 0
    for index, sheet in enumerate(workbook.worksheets):
        sheet.sheet_view.tabSelected = index == 0
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(str(path))
