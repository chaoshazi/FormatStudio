from __future__ import annotations

import datetime as dt
import re
from pathlib import Path
from typing import Any

import openpyxl
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..utils import read_csv_rows, write_csv, write_text


def _safe_sheet_title(name: str, used: set[str] | None = None) -> str:
    title = re.sub(r"[\\/*?:\[\]]", "_", name).strip() or "Sheet1"
    title = title[:31]
    used = used or set()
    candidate = title
    index = 2
    while candidate in used:
        suffix = f"_{index}"
        candidate = title[: 31 - len(suffix)] + suffix
        index += 1
    used.add(candidate)
    return candidate


def _coerce(value: str, all_text: bool) -> Any:
    if all_text:
        return value
    text = value.strip()
    if text == "":
        return None
    if re.fullmatch(r"[-+]?0\d+", text):
        return value
    if re.fullmatch(r"[-+]?\d+", text):
        try:
            return int(text)
        except ValueError:
            return value
    if re.fullmatch(r"[-+]?\d*\.\d+", text):
        try:
            return float(text)
        except ValueError:
            return value
    if re.fullmatch(r"\d{4}[-/]\d{1,2}[-/]\d{1,2}", text):
        for fmt in ("%Y-%m-%d", "%Y/%m/%d"):
            try:
                return dt.datetime.strptime(text, fmt).date()
            except ValueError:
                continue
    return value


def _style_header(worksheet, columns: int, fill: str = "DBEAFE", color: str = "1E3A8A") -> None:
    thin = Side(style="thin", color="CBD5E1")
    for col in range(1, columns + 1):
        cell = worksheet.cell(row=1, column=col)
        cell.font = Font(name="微软雅黑", bold=True, color=color, size=11)
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=thin)


def _auto_width(worksheet) -> None:
    for column in worksheet.columns:
        letter = get_column_letter(column[0].column)
        width = 8
        for cell in column[:80]:
            value = "" if cell.value is None else str(cell.value)
            width = max(width, min(50, sum(2 if ord(ch) > 127 else 1 for ch in value) + 2))
        worksheet.column_dimensions[letter].width = width


def csv_to_xlsx(source: Path, destination: Path, all_text: bool = False) -> None:
    headers, rows, _ = read_csv_rows(source)
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = _safe_sheet_title(source.stem)
    for index, header in enumerate(headers, start=1):
        worksheet.cell(row=1, column=index, value=header)
    for row_index, record in enumerate(rows, start=2):
        for col_index, header in enumerate(headers, start=1):
            worksheet.cell(row=row_index, column=col_index, value=_coerce(record.get(header, ""), all_text))
    _style_header(worksheet, max(1, len(headers)))
    worksheet.freeze_panes = "A2"
    _auto_width(worksheet)
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(str(destination))


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        return value.date().isoformat() if value.time() == dt.time(0, 0) else value.isoformat(sep=" ")
    if isinstance(value, (dt.date, dt.time)):
        return value.isoformat()
    return str(value)
def _load_matrix(path: Path, sheet_name: str | None = None) -> tuple[str, list[list[Any]]]:
    workbook = openpyxl.load_workbook(path, data_only=True)
    try:
        worksheet = workbook[sheet_name] if sheet_name and sheet_name in workbook.sheetnames else workbook.worksheets[0]
        rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
        return worksheet.title, rows
    finally:
        workbook.close()


def xlsx_to_csv(source: Path, destination: Path, sheet_name: str | None = None) -> str:
    title, rows = _load_matrix(source, sheet_name)
    rows = [row for row in rows if any(value is not None and str(value).strip() for value in row)]
    if not rows:
        write_csv(destination, [], [])
        return title
    width = max(len(row) for row in rows)
    normalized = []
    for row in rows:
        values = list(row) + [None] * (width - len(row))
        normalized.append([_cell_text(value) for value in values])
    write_csv(destination, normalized[0], normalized[1:])
    return title


def _docx_add_table(document, rows: list[list[Any]], header: bool = True) -> None:
    if not rows:
        return
    columns = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=columns)
    table.style = "Table Grid"
    for row_index, row in enumerate(rows):
        for col_index in range(columns):
            value = row[col_index] if col_index < len(row) else None
            cell = table.cell(row_index, col_index)
            cell.text = ""
            paragraph = cell.paragraphs[0]
            run = paragraph.add_run(_cell_text(value))
            run.font.size = Pt(9)
            if header and row_index == 0:
                run.bold = True
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _prepare_docx(landscape: bool = False):
    document = Document()
    section = document.sections[0]
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    return document


def csv_to_docx(source: Path, destination: Path) -> None:
    headers, records, _ = read_csv_rows(source)
    rows = [headers] + [[record.get(header, "") for header in headers] for record in records]
    document = _prepare_docx(landscape=len(headers) > 6)
    document.add_heading(source.stem, level=1)
    _docx_add_table(document, rows)
    document.save(str(destination))


def xlsx_to_docx(source: Path, destination: Path, sheet_name: str | None = None) -> None:
    workbook = openpyxl.load_workbook(source, data_only=True)
    try:
        sheets = [workbook[sheet_name]] if sheet_name and sheet_name in workbook.sheetnames else workbook.worksheets
        document = _prepare_docx(landscape=any(sheet.max_column > 6 for sheet in sheets))
        for sheet in sheets:
            rows = [list(row) for row in sheet.iter_rows(values_only=True)]
            rows = [row for row in rows if any(value is not None and str(value).strip() for value in row)]
            if not rows:
                continue
            document.add_heading(sheet.title, level=1)
            _docx_add_table(document, rows)
        document.save(str(destination))
    finally:
        workbook.close()
def _markdown_table(rows: list[list[Any]]) -> str:
    if not rows:
        return ""
    columns = max(len(row) for row in rows)
    normalized: list[list[str]] = []
    for row in rows:
        values = [_cell_text(value).replace("\n", "<br>").replace("|", "\\|") for value in row]
        values += [""] * (columns - len(values))
        normalized.append(values)
    lines = ["| " + " | ".join(normalized[0]) + " |", "| " + " | ".join(["---"] * columns) + " |"]
    for row in normalized[1:]:
        lines.append("| " + " | ".join(row) + " |")
    return "\n".join(lines)


def csv_to_md(source: Path, destination: Path) -> None:
    headers, records, _ = read_csv_rows(source)
    rows = [headers] + [[record.get(header, "") for header in headers] for record in records]
    write_text(destination, f"# {source.stem}\n\n{_markdown_table(rows)}\n")


def xlsx_to_md(source: Path, destination: Path, sheet_name: str | None = None) -> None:
    workbook = openpyxl.load_workbook(source, data_only=True)
    try:
        sheets = [workbook[sheet_name]] if sheet_name and sheet_name in workbook.sheetnames else workbook.worksheets
        parts = [f"# {source.stem}"]
        for sheet in sheets:
            rows = [list(row) for row in sheet.iter_rows(values_only=True)]
            rows = [row for row in rows if any(value is not None and str(value).strip() for value in row)]
            if not rows:
                continue
            parts.append(f"## {sheet.title}")
            parts.append(_markdown_table(rows))
        write_text(destination, "\n\n".join(parts) + "\n")
    finally:
        workbook.close()


def docx_to_xlsx(source: Path, destination: Path) -> int:
    document = Document(str(source))
    if not document.tables:
        raise ValueError("文档中没有可提取的表格")
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    used: set[str] = set()
    for index, table in enumerate(document.tables, start=1):
        worksheet = workbook.create_sheet(_safe_sheet_title(f"Table{index}", used))
        for row_index, row in enumerate(table.rows, start=1):
            for col_index, cell in enumerate(row.cells, start=1):
                worksheet.cell(row=row_index, column=col_index, value=cell.text)
        _style_header(worksheet, max(1, len(table.columns)))
        worksheet.freeze_panes = "A2"
        _auto_width(worksheet)
    workbook.save(str(destination))
    return len(document.tables)